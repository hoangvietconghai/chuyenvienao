# -*- coding: utf-8 -*-
"""
Kiểm thử trọn vẹn quy trình WF4 với DeepSeek API:
- Đọc tệp Word dự thảo Kế hoạch 19
- Đọc tệp PDF Tờ trình của UBND xã
- Chạy phân tích thẩm định qua DeepSeek API
- Xuất tệp Báo cáo thẩm định (-BC/VPĐU) và Công văn lấy ý kiến (-CV/ĐU)
- Kiểm tra chi tiết chất lượng văn bản sinh ra
"""

import asyncio
import os
import sys
from pathlib import Path
import docx
import pypdf

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.abspath("."))

from server.workflows.wf4_appraisal import analyze_agency_draft, generate_appraisal_documents


async def run_test():
    draft_docx_path = r"e:\Viet Design\Chuyên viên ảo\ĐU-KH triển khai thực hiện Nghị quyết số 19 của BCH Trung ương Đảng khoá XIV.docx"
    ttr_pdf_path = r"e:\Viet Design\Chuyên viên ảo\van_ban_den\2026\du_thao_co_quan\H32.227-TTr-TTr-0087-2026_dadongdau.pdf"

    print(f"1. Đang đọc tệp Word dự thảo: {draft_docx_path}")
    doc = docx.Document(draft_docx_path)
    draft_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

    print(f"2. Đang đọc tệp PDF Tờ trình: {ttr_pdf_path}")
    reader = pypdf.PdfReader(ttr_pdf_path)
    ttr_text = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])

    print("3. Đang gọi DeepSeek API để thẩm định 2 tầng...")
    analysis_data = await analyze_agency_draft(
        draft_text=draft_text,
        submitting_agency="Uỷ ban nhân dân xã",
        attached_submission=ttr_text
    )

    print("\n--- DỮ LIỆU PHÂN TÍCH NHẬN ĐƯỢC TỪ DEEPSEEK ---")
    print(f"Tên văn bản: {analysis_data.get('thong_tin_du_thao', {}).get('ten_van_ban')}")
    print(f"Số cảnh báo: {len(analysis_data.get('canh_bao_so_lieu', []))}")
    for cb in analysis_data.get('canh_bao_so_lieu', []):
        print(f"  + [{cb.get('muc_van_ban')}]: {cb.get('van_de')} -> Đề xuất: {cb.get('de_xuat_hoi')}")

    print("\n4. Đang xuất tài liệu Word...")
    output_files = generate_appraisal_documents(
        analysis_data=analysis_data,
        source_docx_path=draft_docx_path,
        custom_dir=r"e:\Viet Design\Chuyên viên ảo\van_ban_du_thao"
    )

    print("\n--- CÁC TỆP ĐÃ XUẤT RA ---")
    for k, v in output_files.items():
        print(f"{k}: {v}")

    bc_path = output_files.get("bc_tham_dinh")
    if bc_path and os.path.exists(bc_path):
        print(f"\n5. Kiểm tra tệp Báo cáo thẩm định: {bc_path}")
        bc_doc = docx.Document(bc_path)
        print("Nội dung các đoạn văn bản trong Báo cáo thẩm định:")
        for idx, p in enumerate(bc_doc.paragraphs):
            if p.text.strip():
                print(f"  [{idx}] {p.text}")

        # Kiểm tra bảng Header và Nơi nhận/Ký
        for t_idx, table in enumerate(bc_doc.tables):
            print(f"\n--- Bảng {t_idx+1} ---")
            for row in table.rows:
                cells_text = [c.text.strip().replace("\n", " | ") for c in row.cells]
                print(f"  {' [COT] '.join(cells_text)}")

if __name__ == "__main__":
    asyncio.run(run_test())
