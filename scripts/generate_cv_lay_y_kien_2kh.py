# -*- coding: utf-8 -*-
"""
scripts/generate_cv_lay_y_kien_2kh.py
Tạo 02 tệp Word (.docx) chuẩn thể thức khối Đảng phục vụ lấy ý kiến Ban Thường vụ Đảng uỷ:
1. Công văn lấy ý kiến Ủy viên Ban Thường vụ Đảng uỷ đối với 02 dự thảo Kế hoạch (-CV/ĐU)
2. Phiếu xin ý kiến Ủy viên Ban Thường vụ Đảng uỷ về 02 dự thảo Kế hoạch
"""

import os
import sys
from pathlib import Path

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT

try:
    from scripts.export_docx import (
        PartyDocumentBuilder,
        set_cell_margins,
        remove_table_borders,
        add_horizontal_line,
        parse_markdown_runs
    )
except ImportError:
    from export_docx import (
        PartyDocumentBuilder,
        set_cell_margins,
        remove_table_borders,
        add_horizontal_line,
        parse_markdown_runs
    )


def generate_cv_lay_y_kien_2kh():
    """Tạo Công văn lấy ý kiến 02 dự thảo Kế hoạch chung trong 01 công văn."""
    data_cv = {
        "doc_type": "CV",
        "so_hieu": "Số        -CV/ĐU",
        "co_quan_cap_tren": "ĐẢNG BỘ TỈNH KHÁNH HOÀ",
        "co_quan_ban_hanh": "ĐẢNG UỶ XÃ CÔNG HẢI",
        "dia_danh_ngay": "Công Hải, ngày   tháng 10 năm 2026",
        "trich_yeu": "Góp ý các Dự thảo Kế hoạch của\nBan Thường vụ Đảng uỷ",
        "kinh_gui": ["Các đồng chí Uỷ viên Ban Thường vụ Đảng uỷ."],
        "noi_dung": [
            "Tiếp nhận các Tờ trình của Uỷ ban nhân dân xã về việc xem xét, cho ý kiến đối với các dự thảo Kế hoạch triển khai thực hiện các văn bản chỉ đạo của Tỉnh uỷ Khánh Hoà và Trung ương được Ban Thường vụ Đảng uỷ giao; Thường trực Đảng uỷ kính gửi các đồng chí Uỷ viên Ban Thường vụ Đảng uỷ nghiên cứu, rà soát, đề xuất và tham gia đóng góp ý kiến đối với 02 Dự thảo Kế hoạch của Ban Thường vụ Đảng uỷ, cụ thể gồm:",
            "- Dự thảo Kế hoạch triển khai thực hiện Kế hoạch số 106-KH/TU ngày 28/8/2026 của Ban Chấp hành Đảng bộ tỉnh về thực hiện Nghị quyết số 20-NQ/TW ngày 28/7/2026 của Ban Chấp hành Trung ương Đảng khoá XIV và Chương trình hành động số 34-CTr/TW ngày 28/7/2026 của Bộ Chính trị về xây dựng và phát triển Việt Nam trở thành quốc gia biển mạnh trên địa bàn xã Công Hải.",
            "- Dự thảo Kế hoạch thực hiện Chương trình hành động số 30-CTr/TU ngày 28/8/2026 của Tỉnh uỷ về thực hiện Nghị quyết số 19-NQ/TW của Ban Chấp hành Trung ương Đảng khoá XIV và Chương trình hành động số 33-CTr/TW ngày 28/7/2026 của Bộ Chính trị về đổi mới mô hình phát triển Việt Nam trên địa bàn xã Công Hải.",
            "*(Hồ sơ gửi kèm gồm: 02 bản dự thảo Kế hoạch đã được Văn phòng Đảng uỷ phối hợp với cơ quan chủ trì rà soát, chuẩn hoá thể thức theo quy định và Phiếu xin ý kiến)*.",
            "Đề nghị các đồng chí nghiên cứu, tham gia đóng góp ý kiến bằng văn bản (theo mẫu Phiếu xin ý kiến gửi kèm) hoặc góp ý trực tiếp vào các bản dự thảo (gửi kèm); nội dung góp ý gửi về Thường trực Đảng uỷ (qua Văn phòng Đảng uỷ) **trước ngày 12/10/2026** để tổng hợp, tiếp thu và hoàn thiện văn bản trình đồng chí Bí thư Đảng uỷ ký ban hành chính thức.",
            "Qua thời gian nêu trên, nếu các đồng chí không có ý kiến gửi về xem như thống nhất với nội dung các Dự thảo Kế hoạch./."
        ],
        "noi_nhan": [
            "- Như trên;",
            "- Thường trực Đảng uỷ (để b/c);",
            "- Lưu VPĐU (CV-2026)."
        ],
        "tham_quyen": "T/L BAN THƯỜNG VỤ",
        "chuc_vu": "PHÓ CHÁNH VĂN PHÒNG",
        "nguoi_ky": "Ngô Hoàng Việt"
    }

    builder = PartyDocumentBuilder(data_cv)
    builder.build_header()
    builder.build_title_section()
    builder.build_body()
    builder.build_footer()

    out_dir = BASE_DIR / "van_ban_du_thao" / "2026" / "Cong_van"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = str(out_dir / "CV_lay_y_kien_BTV_ve_02_du_thao_Ke_hoach_2026.docx")
    builder.save(out_file)

    out_file_root = str(BASE_DIR / "van_ban_du_thao" / "CV_lay_y_kien_BTV_ve_02_du_thao_Ke_hoach_2026.docx")
    builder.save(out_file_root)

    print(f"[OK] Đã xuất Công văn lấy ý kiến: {out_file}")
    return out_file


def generate_phieu_xin_y_kien_2kh():
    """Tạo tệp Phiếu xin ý kiến Ủy viên Ban Thường vụ Đảng uỷ đối với 02 dự thảo Kế hoạch."""
    doc = docx.Document()

    # Khổ giấy A4, Lề trang chuẩn: Trái 30mm, Phải 15mm, Trên 20mm, Dưới 20mm
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)

    # Style
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    style.paragraph_format.line_spacing = Pt(18)
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(6)

    # 1. Header Table
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(tbl)

    col1 = tbl.columns[0]
    col2 = tbl.columns[1]
    col1.width = Mm(75)
    col2.width = Mm(90)
    cell_l = tbl.cell(0, 0)
    cell_r = tbl.cell(0, 1)
    cell_l.width = Mm(75)
    cell_r.width = Mm(90)
    set_cell_margins(cell_l, top=0, bottom=0, left=0, right=70)
    set_cell_margins(cell_r, top=0, bottom=0, left=70, right=0)

    # Cột trái
    p1 = cell_l.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p1.paragraph_format.line_spacing = 1.05
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(1)
    r = p1.add_run("ĐẢNG BỘ TỈNH KHÁNH HOÀ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    p2 = cell_l.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p2.paragraph_format.line_spacing = 1.05
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(1)
    r = p2.add_run("ĐẢNG UỶ XÃ CÔNG HẢI")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True

    p3 = cell_l.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p3.paragraph_format.line_spacing = 1.0
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(2)
    r = p3.add_run("*")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True

    # Cột phải
    p_tn = cell_r.paragraphs[0]
    p_tn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tn.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_tn.paragraph_format.line_spacing = 1.05
    p_tn.paragraph_format.space_before = Pt(0)
    p_tn.paragraph_format.space_after = Pt(1)
    r = p_tn.add_run("ĐẢNG CỘNG SẢN VIỆT NAM")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True

    p_line = cell_r.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_line.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_line.paragraph_format.line_spacing = 1.0
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(4)
    add_horizontal_line(p_line, width_pt=192, weight_pt=0.75)

    p_ngay = cell_r.add_paragraph()
    p_ngay.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ngay.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_ngay.paragraph_format.line_spacing = 1.05
    p_ngay.paragraph_format.space_before = Pt(0)
    p_ngay.paragraph_format.space_after = Pt(0)
    r = p_ngay.add_run("Công Hải, ngày   tháng 10 năm 2026")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.italic = True

    # Khoảng cách
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(6)
    p_sp.paragraph_format.space_after = Pt(4)

    # 2. Tiêu đề Phiếu
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_title.paragraph_format.line_spacing = Pt(18)
    r = p_title.add_run("PHIẾU XIN Ý KIẾN")
    r.font.name = "Times New Roman"
    r.font.size = Pt(15)
    r.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    p_sub.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_sub.paragraph_format.line_spacing = Pt(18)
    r = p_sub.add_run("CỦA ĐỒNG CHÍ UỶ VIÊN BAN THƯỜNG VỤ ĐẢNG UỶ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True

    p_trich = doc.add_paragraph()
    p_trich.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_trich.paragraph_format.space_before = Pt(0)
    p_trich.paragraph_format.space_after = Pt(2)
    p_trich.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_trich.paragraph_format.line_spacing = Pt(18)
    r = p_trich.add_run("Về 02 dự thảo Kế hoạch của Ban Thường vụ Đảng uỷ:\n(1) Triển khai Kế hoạch số 106-KH/TU của Tỉnh uỷ; (2) Triển khai CTr số 30-CTr/TU của Tỉnh uỷ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13.5)
    r.font.bold = True

    p_kem = doc.add_paragraph()
    p_kem.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kem.paragraph_format.space_before = Pt(0)
    p_kem.paragraph_format.space_after = Pt(10)
    p_kem.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_kem.paragraph_format.line_spacing = Pt(18)
    r = p_kem.add_run("(Kèm theo Công văn số        -CV/ĐU ngày   /10/2026 của Ban Thường vụ Đảng uỷ)")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.italic = True

    # 3. Kính gửi
    p_kg = doc.add_paragraph()
    p_kg.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kg.paragraph_format.first_line_indent = Mm(10)
    p_kg.paragraph_format.space_before = Pt(6)
    p_kg.paragraph_format.space_after = Pt(6)
    p_kg.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_kg.paragraph_format.line_spacing = Pt(18)
    r = p_kg.add_run("Kính gửi: ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.italic = True
    r2 = p_kg.add_run("Đồng chí ......................................................, Uỷ viên Ban Thường vụ Đảng uỷ xã.")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)

    # 4. Thân phiếu
    p_dan = doc.add_paragraph()
    p_dan.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_dan.paragraph_format.first_line_indent = Mm(10)
    p_dan.paragraph_format.space_before = Pt(6)
    p_dan.paragraph_format.space_after = Pt(6)
    p_dan.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_dan.paragraph_format.line_spacing = Pt(18)
    r = p_dan.add_run("Thực hiện Quy chế làm việc của Ban Chấp hành Đảng bộ xã khoá I, nhiệm kỳ 2025 - 2030, Thường trực Đảng uỷ xin ý kiến đồng chí về 02 dự thảo Kế hoạch của Ban Thường vụ Đảng uỷ xã như sau:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    # Mục I: Kế hoạch 1
    p_kh1 = doc.add_paragraph()
    p_kh1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kh1.paragraph_format.first_line_indent = Mm(10)
    p_kh1.paragraph_format.space_before = Pt(6)
    p_kh1.paragraph_format.space_after = Pt(4)
    p_kh1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_kh1.paragraph_format.line_spacing = Pt(18)
    parse_markdown_runs(p_kh1, "**I. Đối với Dự thảo Kế hoạch triển khai Kế hoạch số 106-KH/TU của Tỉnh uỷ (về quốc gia biển mạnh):**", base_font_size=14, default_bold=True)

    opt_list_1 = [
        "[   ] Thống nhất hoàn toàn với nội dung dự thảo Kế hoạch.",
        "[   ] Cơ bản thống nhất, có ý kiến tham gia cụ thể: ........................................................................",
        ".........................................................................................................................................................................",
        "[   ] Không thống nhất với dự thảo (nêu rõ lý do): ..............................................................................",
        "........................................................................................................................................................................."
    ]
    for opt in opt_list_1:
        p_opt = doc.add_paragraph()
        p_opt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_opt.paragraph_format.first_line_indent = Mm(15)
        p_opt.paragraph_format.space_before = Pt(2)
        p_opt.paragraph_format.space_after = Pt(2)
        p_opt.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p_opt.paragraph_format.line_spacing = Pt(18)
        r = p_opt.add_run(opt)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)

    # Mục II: Kế hoạch 2
    p_kh2 = doc.add_paragraph()
    p_kh2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kh2.paragraph_format.first_line_indent = Mm(10)
    p_kh2.paragraph_format.space_before = Pt(6)
    p_kh2.paragraph_format.space_after = Pt(4)
    p_kh2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_kh2.paragraph_format.line_spacing = Pt(18)
    parse_markdown_runs(p_kh2, "**II. Đối với Dự thảo Kế hoạch thực hiện Chương trình hành động số 30-CTr/TU của Tỉnh uỷ (về đổi mới mô hình phát triển):**", base_font_size=14, default_bold=True)

    opt_list_2 = [
        "[   ] Thống nhất hoàn toàn với nội dung dự thảo Kế hoạch.",
        "[   ] Cơ bản thống nhất, có ý kiến tham gia cụ thể: ........................................................................",
        ".........................................................................................................................................................................",
        "[   ] Không thống nhất với dự thảo (nêu rõ lý do): ..............................................................................",
        "........................................................................................................................................................................."
    ]
    for opt in opt_list_2:
        p_opt = doc.add_paragraph()
        p_opt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_opt.paragraph_format.first_line_indent = Mm(15)
        p_opt.paragraph_format.space_before = Pt(2)
        p_opt.paragraph_format.space_after = Pt(2)
        p_opt.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p_opt.paragraph_format.line_spacing = Pt(18)
        r = p_opt.add_run(opt)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)

    # Mục III: Ý kiến khác
    p_kh3 = doc.add_paragraph()
    p_kh3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kh3.paragraph_format.first_line_indent = Mm(10)
    p_kh3.paragraph_format.space_before = Pt(6)
    p_kh3.paragraph_format.space_after = Pt(4)
    p_kh3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_kh3.paragraph_format.line_spacing = Pt(18)
    parse_markdown_runs(p_kh3, "**III. Ý kiến đề xuất khác (nếu có):**", base_font_size=14, default_bold=True)

    dots = [
        ".........................................................................................................................................................................",
        "........................................................................................................................................................................."
    ]
    for dot in dots:
        p_d = doc.add_paragraph()
        p_d.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_d.paragraph_format.first_line_indent = Mm(10)
        p_d.paragraph_format.space_before = Pt(2)
        p_d.paragraph_format.space_after = Pt(2)
        p_d.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p_d.paragraph_format.line_spacing = Pt(18)
        r = p_d.add_run(dot)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)

    # Ghi chú
    p_note = doc.add_paragraph()
    p_note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_note.paragraph_format.first_line_indent = Mm(10)
    p_note.paragraph_format.space_before = Pt(8)
    p_note.paragraph_format.space_after = Pt(12)
    p_note.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_note.paragraph_format.line_spacing = Pt(18)
    r = p_note.add_run("* Phiếu xin ý kiến đề nghị gửi về Thường trực Đảng uỷ (qua Văn phòng Đảng uỷ) trước ngày 12/10/2026 để tổng hợp báo cáo Thường trực Đảng uỷ xem xét, hoàn thiện văn bản ban hành.*")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.italic = True

    # Khối ký tên của UVBTV
    p_sign_t = doc.add_paragraph()
    p_sign_t.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign_t.paragraph_format.right_indent = Mm(10)
    p_sign_t.paragraph_format.space_before = Pt(6)
    p_sign_t.paragraph_format.space_after = Pt(2)
    p_sign_t.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_sign_t.paragraph_format.line_spacing = 1.05
    r = p_sign_t.add_run("UỶ VIÊN BAN THƯỜNG VỤ ĐẢNG UỶ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13.5)
    r.font.bold = True

    p_sign_sub = doc.add_paragraph()
    p_sign_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign_sub.paragraph_format.right_indent = Mm(25)
    p_sign_sub.paragraph_format.space_before = Pt(0)
    p_sign_sub.paragraph_format.space_after = Pt(0)
    p_sign_sub.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_sign_sub.paragraph_format.line_spacing = 1.05
    r = p_sign_sub.add_run("(Ký và ghi rõ họ tên)")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.italic = True

    for _ in range(5):
        p_blank = doc.add_paragraph()
        p_blank.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p_blank.paragraph_format.line_spacing = Pt(14)
        p_blank.paragraph_format.space_before = Pt(0)
        p_blank.paragraph_format.space_after = Pt(0)

    out_dir = BASE_DIR / "van_ban_du_thao" / "2026" / "Cong_van"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = str(out_dir / "Phieu_xin_y_kien_BTV_ve_02_du_thao_Ke_hoach_2026.docx")
    doc.save(out_file)

    out_file_root = str(BASE_DIR / "van_ban_du_thao" / "Phieu_xin_y_kien_BTV_ve_02_du_thao_Ke_hoach_2026.docx")
    doc.save(out_file_root)

    print(f"[OK] Đã xuất Phiếu xin ý kiến: {out_file}")
    return out_file


if __name__ == "__main__":
    generate_cv_lay_y_kien_2kh()
    generate_phieu_xin_y_kien_2kh()
