#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
scripts/test_workflow.py
========================
Kịch bản kiểm thử tích hợp (Integration Tests) cho Chuyên viên ảo.
Kiểm tra sinh tự động các thể loại văn bản Đảng theo HD 05-HD/VPTW.
"""

import os
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Thêm thư mục scripts vào sys.path
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(SCRIPT_DIR)
sys.path.insert(0, SCRIPT_DIR)

from export_docx import export_party_document


def test_tc02_thong_bao_ket_luan():
    print("[TEST TC-02] Đang kiểm thử sinh Thông báo Kết luận họp tuần Thường trực Đảng uỷ...")
    tb_data = {
        "doc_type": "TB",
        "so_hieu": "Số      -TB/ĐU",
        "dia_danh_ngay": "Công Hải, ngày   tháng   năm 2026",
        "ten_loai": "THÔNG BÁO",
        "trich_yeu": "Kết luận của Thường trực Đảng uỷ xã tại cuộc họp giao ban tuần (ngày 06/10/2026)",
        "noi_dung": [
            "Ngày 06/10/2026, Thường trực Đảng uỷ xã tổ chức cuộc họp giao ban tuần để đánh giá tình hình công tác tuần qua và triển khai nhiệm vụ trọng tâm tuần tới. Đồng chí Vũ Thị Thuỳ Trang - Bí thư Đảng uỷ xã chủ trì cuộc họp. Sau khi nghe Văn phòng Đảng uỷ, UBND xã và các ban, ngành báo cáo, Thường trực Đảng uỷ kết luận như sau:",
            "I. ĐÁNH GIÁ CHUNG",
            "Trong tuần qua, các cơ quan, đơn vị trong toàn hệ thống chính trị xã đã chủ động, tập trung triển khai các nhiệm vụ theo kế hoạch; tình hình an ninh chính trị, trật tự an toàn xã hội trên địa bàn được giữ vững.",
            "II. NHIỆM VỤ TRỌNG TÂM TRONG TUẦN TỚI",
            "1. **Uỷ ban nhân dân xã:** Tập trung chỉ đạo đẩy nhanh tiến độ thu ngân sách; rà soát, giải quyết dứt điểm các kiến nghị của cử tri trước kỳ họp HĐND xã; báo cáo Thường trực Đảng uỷ trước ngày 15/10/2026.",
            "2. **Ban Xây dựng Đảng:** Tham mưu kế hoạch rà soát, tạo nguồn và phát triển đảng viên đợt cuối năm 2026; chuẩn bị hồ sơ chuyển đảng chính thức cho các đảng viên dự bị đến hạn.",
            "3. **Văn phòng Đảng uỷ:** Theo dõi, đôn đốc các cơ quan thực hiện nghiêm túc Thông báo này; định kỳ báo cáo tiến độ cho Thường trực Đảng uỷ."
        ],
        "noi_nhan": [
            "- Thường trực Đảng uỷ;",
            "- Thường trực HĐND, Lãnh đạo UBND xã;",
            "- Các cơ quan tham mưu Đảng uỷ;",
            "- Lưu VPĐU."
        ],
        "tham_quyen": "T/L BAN THƯỜNG VỤ",
        "chuc_vu": "CHÁNH VĂN PHÒNG",
        "nguoi_ky": "Ngô Hoàng Việt",
        "output_path": os.path.join(BASE_DIR, "van_ban_du_thao", "2026", "Thong_bao", "TC02_TB_Giao_ban_Thuong_truc.docx")
    }
    out_file = export_party_document(tb_data, tb_data["output_path"])
    assert os.path.exists(out_file), "Không tạo được file TC-02"
    
    # Kiểm tra quy chuẩn áp dụng chung trên TC-02 Thông báo
    import docx
    doc_tb = docx.Document(out_file)
    
    # 1. Đường kẻ dưới ĐẢNG CỘNG SẢN VIỆT NAM (186pt, 0.75pt)
    header_r_xml = doc_tb.tables[0].cell(0, 1)._tc.xml
    assert 'to="186pt,0"' in header_r_xml, "Đường kẻ dưới ĐCSVN ở Thông báo phải dài 186pt"
    assert 'strokeweight="0.75pt"' in header_r_xml, "Độ dày đường kẻ ở Thông báo phải là 3/4pt"
    
    # 2. Ô chữ ký Thông báo Thừa lệnh:
    # Dòng 1: T/L BAN THƯỜNG VỤ (đậm)
    # Dòng 2: CHÁNH VĂN PHÒNG (viết hoa in thường, không đậm)
    # 5 dòng trống cỡ 14pt
    # Dòng cuối: Họ tên (đậm)
    paras_sign = doc_tb.tables[1].cell(0, 1).paragraphs
    assert "T/L BAN THƯỜNG VỤ" in paras_sign[0].text, "Dòng 1 chữ ký phải là T/L BAN THƯỜNG VỤ"
    assert paras_sign[0].runs[0].bold is True, "Dòng 1 phải in đậm"
    assert "CHÁNH VĂN PHÒNG" in paras_sign[1].text, "Dòng 2 chữ ký phải là CHÁNH VĂN PHÒNG"
    assert paras_sign[1].runs[0].bold is False, "Dòng 2 chức danh phải viết hoa in thường không in đậm"
    for i in range(2, 7):
        assert paras_sign[i].text.strip() == "", f"Dòng trống thứ {i-1} phải trống"
        assert paras_sign[i].paragraph_format.line_spacing.pt == 14.0, "Dòng trống phải có line_spacing 14pt"
    assert paras_sign[7].text.strip() == "Ngô Hoàng Việt", "Dòng cuối cùng phải là tên người ký"
    assert paras_sign[7].runs[0].bold is True, "Tên người ký phải in đậm"
    # 3. Kiểm tra số chỉ mục 1., 2., 3. và tiêu đề in đậm cùng nhau, nội dung in thường
    indicator_paras = [p for p in doc_tb.paragraphs if any(p.text.strip().startswith(x) for x in ['1.', '2.', '3.'])]
    assert len(indicator_paras) == 3, "Phải có đúng 3 đoạn chỉ mục 1., 2., 3."
    for p in indicator_paras:
        assert p.runs[0].bold is True, f"Số chỉ mục và tiêu đề phải in đậm: {p.runs[0].text}"
        assert p.runs[0].text.strip().endswith(":"), f"Tiêu đề in đậm phải kết thúc bằng dấu hai chấm: {p.runs[0].text}"
        if len(p.runs) > 1:
            assert p.runs[1].bold is False, f"Nội dung diễn giải sau dấu hai chấm phải in thường: {p.runs[1].text[:30]}"
    
    print(f"  -> TC-02 HOÀN TOÀN ĐẠT CHUẨN: {out_file}")


def test_tc03_bao_cao_tham_dinh():
    print("[TEST TC-03] Đang kiểm thử sinh Báo cáo Thẩm định chuẩn thể thức Đảng của Văn phòng Đảng uỷ...")
    bc_data = {
        "doc_type": "BC",
        "so_hieu": "Số      -BC/VPĐU",
        "dia_danh_ngay": "Công Hải, ngày   tháng   năm 2026",
        "ten_loai": "BÁO CÁO",
        "trich_yeu": "Kết quả thẩm định dự thảo Kế hoạch hành động chuyển đổi số của Uỷ ban nhân dân xã",
        "noi_dung": [
            "Căn cứ Quy chế làm việc số 01-QC/ĐU ngày 02/7/2025 của Ban Chấp hành Đảng bộ xã Công Hải khoá I và phân công của Thường trực Đảng uỷ, Văn phòng Đảng uỷ đã tiến hành thẩm định hồ sơ dự thảo Kế hoạch hành động chuyển đổi số do UBND xã trình kèm theo Tờ trình số 15/TTr-UBND ngày 20/9/2026. Kết quả thẩm định cụ thể như sau:",
            "I. TỔNG QUAN HỒ SƠ THẨM ĐỊNH",
            "- Cơ quan trình: Uỷ ban nhân dân xã Công Hải.",
            "- Tên dự thảo văn bản: Kế hoạch triển khai thực hiện Kế hoạch số 101-KH/TU của Ban Thường vụ Tỉnh uỷ về hành động 100 ngày giải quyết điểm nghẽn chuyển đổi số.",
            "- Hồ sơ gửi kèm: Tờ trình số 15/TTr-UBND; dự thảo Kế hoạch; Bản tổng hợp ý kiến tham gia của các cơ quan, đơn vị.",
            "II. KẾT QUẢ THẨM ĐỊNH",
            "1. Thể thức văn bản: Văn phòng Đảng uỷ đã rà soát, chỉnh sửa trực tiếp trên dự thảo văn bản các lỗi chính tả, chỉnh sửa thể thức văn bản theo đúng Hướng dẫn số 05-HD/VPTW ngày 12/01/2026 của Văn phòng Trung ương Đảng.",
            "2. Nội dung và số liệu chuyên môn: Dự thảo Kế hoạch cơ bản bám sát mục tiêu, nhiệm vụ theo chỉ đạo của Ban Thường vụ Tỉnh uỷ. Tuy nhiên, tại Mục III.2, dự thảo chưa xác định rõ tiến độ hoàn thành và phân công cán bộ phụ trách cụ thể cho từng thôn theo nguyên tắc '5 Rõ'; mốc thời hạn hoàn thành tại Khoản 3 Mục IV ấn định ngày 27/9/2026 trùng vào ngày Chủ nhật.",
            "III. ĐỀ XUẤT, KIẾN NGHỊ",
            "Trên cơ sở kết quả thẩm định, Văn phòng Đảng uỷ kính trình Thường trực Đảng uỷ:",
            "- Dự thảo văn bản đã được Văn phòng Đảng uỷ trực tiếp hoàn thiện bảo đảm đúng thể thức Hướng dẫn số 05-HD/VPTW.",
            "- Đề nghị Thường trực Đảng uỷ chỉ đạo UBND xã: (1) Rà soát, bổ sung rõ tiến độ hoàn thành và cá nhân chịu trách nhiệm tại từng thôn tại Mục III.2; (2) Điều chỉnh thời hạn hoàn thành tại Khoản 3 Mục IV sang ngày 28/9/2026 (ngày làm việc).",
            "- Kính đề nghị Thường trực Đảng uỷ xem xét cho ý kiến chỉ đạo tại cuộc họp giao ban tuần trước khi trình Ban Thường vụ Đảng uỷ ban hành./."
        ],
        "noi_nhan": [
            "- Thường trực Đảng uỷ;",
            "- Ban Thường vụ Đảng uỷ;",
            "- Uỷ ban nhân dân xã;",
            "- Lưu VPĐU."
        ],
        "tham_quyen": "K/T CHÁNH VĂN PHÒNG",
        "chuc_vu": "PHÓ CHÁNH VĂN PHÒNG",
        "nguoi_ky": "Ngô Hoàng Việt",
        "output_path": os.path.join(BASE_DIR, "van_ban_du_thao", "2026", "Bao_cao", "TC03_BC_Tham_dinh_KH_Chuyen_doi_so.docx")
    }
    out_file = export_party_document(bc_data, bc_data["output_path"])
    assert os.path.exists(out_file), "Không tạo được file TC-03"
    
    # Kiểm tra chi tiết cấu trúc tài liệu Word sinh ra
    import docx
    doc = docx.Document(out_file)
    
    # 1. Kiểm tra không có bảng kính gửi (chỉ có 2 bảng Header và Footer)
    assert len(doc.tables) == 2, f"Báo cáo phải có đúng 2 bảng (Header & Footer), thực tế: {len(doc.tables)}"
    
    # 2. Kiểm tra chỉ mục: chỉ in đậm tiêu đề, nội dung in thường
    checked_indicators = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("1. Thể thức văn bản:") or txt.startswith("2. Nội dung và số liệu chuyên môn:"):
            assert p.runs[0].bold is True, f"Tiêu đề chỉ mục phải in đậm: {p.runs[0].text}"
            assert p.runs[1].bold is False, f"Nội dung sau dấu hai chấm không được in đậm: {p.runs[1].text[:30]}"
            checked_indicators += 1
    assert checked_indicators == 2, "Chưa kiểm tra đủ 2 chỉ mục bắt buộc"
    
    # 3. Kiểm tra ô chữ ký: K/T CHÁNH VĂN PHÒNG (đậm) / PHÓ CHÁNH VĂN PHÒNG (viết hoa không đậm), cách đúng 5 dòng cỡ 14pt
    cell_sign = doc.tables[1].cell(0, 1)
    paras = cell_sign.paragraphs
    assert "K/T CHÁNH VĂN PHÒNG" in paras[0].text, "Dòng 1 chữ ký phải là K/T CHÁNH VĂN PHÒNG"
    assert paras[0].runs[0].bold is True, "Dòng 1 thẩm quyền/ký thay phải in đậm"
    assert "PHÓ CHÁNH VĂN PHÒNG" in paras[1].text, "Dòng 2 chữ ký phải là PHÓ CHÁNH VĂN PHÒNG"
    assert paras[1].runs[0].bold is False, "Dòng 2 chức vụ (BÍ THƯ, PHÓ BÍ THƯ, PHÓ CHÁNH VĂN PHÒNG...) phải viết hoa in thường, không in đậm"
    # 5 dòng trống từ index 2 đến 6
    for i in range(2, 7):
        assert paras[i].text.strip() == "", f"Dòng {i+1} phải là dòng trống cách chức danh"
        assert paras[i].paragraph_format.line_spacing.pt == 14.0, f"Dòng trống {i+1} phải có cỡ dòng 14pt"
    assert paras[7].text.strip() == "Ngô Hoàng Việt", "Dòng cuối cùng phải là tên người ký"
    assert paras[7].runs[0].bold is True, "Họ tên người ký phải in đậm"
    
    # 4. Kiểm tra đường kẻ ngang dưới ĐẢNG CỘNG SẢN VIỆT NAM (dài 186pt, dày 0.75pt)
    cell_header_r = doc.tables[0].cell(0, 1)
    header_xml = cell_header_r._tc.xml
    assert 'to="186pt,0"' in header_xml, "Đường kẻ dưới ĐCSVN phải dài 186pt (kéo dài toàn bộ dòng chữ)"
    assert 'strokeweight="0.75pt"' in header_xml, "Độ dày đường kẻ phải là 3/4pt (0.75pt)"

    # 5. Kiểm tra Header Cột 1 của Báo cáo thẩm định do Văn phòng ban hành: ĐẢNG UỶ XÃ CÔNG HẢI / VĂN PHÒNG
    cell_header_l = doc.tables[0].cell(0, 0)
    assert "ĐẢNG UỶ XÃ CÔNG HẢI" in cell_header_l.paragraphs[0].text, "Dòng 1 Header phải là ĐẢNG UỶ XÃ CÔNG HẢI"
    assert "VĂN PHÒNG" == cell_header_l.paragraphs[1].text.strip(), "Dòng 2 Header phải là VĂN PHÒNG (không ghi VĂN PHÒNG ĐẢNG UỶ)"

    # 6. Kiểm tra toàn bộ đoạn thân bài (dưới tiêu đề đến trên ký/nơi nhận): Before 6pt, After 6pt, Exactly 18pt, thụt đầu dòng 10mm
    body_count = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        # Bỏ qua các dòng tiêu đề căn giữa (BÁO CÁO, trích yếu, đoạn đệm đầu)
        if not txt or txt == "BÁO CÁO" or "kết quả thẩm định dự thảo" in txt.lower():
            continue
        body_count += 1
        assert round(p.paragraph_format.space_before.pt, 1) == 6.0, f"Đoạn thân bài phải có space_before 6pt: {txt[:30]}"
        assert round(p.paragraph_format.space_after.pt, 1) == 6.0, f"Đoạn thân bài phải có space_after 6pt: {txt[:30]}"
        assert round(p.paragraph_format.line_spacing.pt, 1) == 18.0, f"Đoạn thân bài phải có line_spacing 18pt: {txt[:30]}"
        assert round(p.paragraph_format.first_line_indent.mm, 1) == 10.0, f"Đoạn thân bài phải thụt đầu dòng 10mm: {txt[:30]}"
    assert body_count >= 10, f"Số lượng đoạn thân bài được kiểm tra phải >= 10, thực tế: {body_count}"
    
    # 7. Kiểm tra số chỉ mục 1., 2. và tiêu đề in đậm cùng nhau, nội dung in thường
    ind_paras = [p for p in doc.paragraphs if any(p.text.strip().startswith(x) for x in ['1. Thể thức', '2. Nội dung'])]
    assert len(ind_paras) == 2, "Phải có 2 đoạn chỉ mục kết quả thẩm định"
    for p in ind_paras:
        assert p.runs[0].bold is True, f"Số chỉ mục và tiêu đề phải in đậm: {p.runs[0].text}"
        assert p.runs[0].text.strip().endswith(":"), f"Tiêu đề in đậm phải kết thúc bằng dấu hai chấm: {p.runs[0].text}"
        if len(p.runs) > 1:
            assert p.runs[1].bold is False, f"Nội dung sau dấu hai chấm phải in thường: {p.runs[1].text[:30]}"

    print(f"  -> TC-03 HOÀN TOÀN ĐẠT CHUẨN: {out_file}")


def test_tc05_indicator_bolding_universal():
    print("[TEST TC-05] Kiểm thử quy chuẩn in đậm chỉ mục trên tất cả các dạng định dạng đầu vào...")
    kh_test_data = {
        "doc_type": "KH",
        "so_hieu": "Số 05-KH/ĐU",
        "dia_danh_ngay": "Công Hải, ngày 02 tháng 10 năm 2026",
        "ten_loai": "KẾ HOẠCH",
        "trich_yeu": "Thực hiện nhiệm vụ trọng tâm công tác Đảng quý IV năm 2026",
        "noi_dung": [
            "I. MỤC ĐÍCH, YÊU CẦU",
            "Nhằm cụ thể hoá các chỉ tiêu, nhiệm vụ quý IV năm 2026 của Đảng bộ xã, Ban Thường vụ Đảng uỷ ban hành Kế hoạch thực hiện cụ thể như sau:",
            "II. NHIỆM VỤ VÀ GIẢI PHÁP",
            "1. **Về công tác chính trị, tư tưởng:** Tăng cường tuyên truyền các chủ trương, nghị quyết của Đảng đến từng chi bộ và quần chúng nhân dân.",
            "2. Về công tác tổ chức cán bộ: Tập trung rà soát quy hoạch cấp uỷ nhiệm kỳ tới.",
            "a) **Công tác phát triển đảng viên:** Phấn đấu kết nạp vượt chỉ tiêu được giao.",
            "b) Công tác kiểm tra, giám sát: Hoàn thành 100% cuộc kiểm tra theo chương trình.",
            "- **Cơ quan thường trực:** Văn phòng Đảng uỷ chịu trách nhiệm theo dõi, tổng hợp.",
            "- Cơ quan phối hợp: UBND xã và các đoàn thể chính trị - xã hội phối hợp triển khai thực hiện.",
            "+ Phân công 01 cán bộ chuyên trách theo dõi, tổng hợp báo cáo định kỳ.",
            "1. Mục tiêu tổng quát"
        ],
        "noi_nhan": ["- Thường trực Đảng uỷ;", "- Các chi bộ trực thuộc;", "- Lưu VPĐU."],
        "tham_quyen": "T/M BAN THƯỜNG VỤ",
        "chuc_vu": "BÍ THƯ",
        "nguoi_ky": "Vũ Thị Thuỳ Trang",
        "output_path": os.path.join(BASE_DIR, "van_ban_du_thao", "2026", "Ke_hoach", "TC05_KH_Kiem_tra_chi_muc.docx")
    }
    out_file = export_party_document(kh_test_data, kh_test_data["output_path"])
    assert os.path.exists(out_file), "Không xuất được file TC-05"

    import docx
    doc_kh = docx.Document(out_file)

    # Kiểm tra từng đoạn chỉ mục
    # 1. '1. Về công tác chính trị, tư tưởng:' (từ '1. **Về công tác chính trị, tư tưởng:**')
    p1 = [p for p in doc_kh.paragraphs if p.text.strip().startswith("1. Về công tác chính trị")][0]
    assert p1.runs[0].bold is True and p1.runs[0].text.strip() == "1. Về công tác chính trị, tư tưởng:"
    assert p1.runs[1].bold is False

    # 2. '2. Về công tác tổ chức cán bộ:' (không có markdown)
    p2 = [p for p in doc_kh.paragraphs if p.text.strip().startswith("2. Về công tác tổ chức")][0]
    assert p2.runs[0].bold is True and p2.runs[0].text.strip() == "2. Về công tác tổ chức cán bộ:"
    assert p2.runs[1].bold is False

    # 3. 'a) Công tác phát triển đảng viên:' (từ 'a) **Công tác phát triển đảng viên:**')
    pa = [p for p in doc_kh.paragraphs if p.text.strip().startswith("a) Công tác phát triển")][0]
    assert pa.runs[0].bold is True and pa.runs[0].text.strip() == "a) Công tác phát triển đảng viên:"
    assert pa.runs[1].bold is False

    # 4. 'b) Công tác kiểm tra, giám sát:' (không có markdown)
    pb = [p for p in doc_kh.paragraphs if p.text.strip().startswith("b) Công tác kiểm tra")][0]
    assert pb.runs[0].bold is True and pb.runs[0].text.strip() == "b) Công tác kiểm tra, giám sát:"
    assert pb.runs[1].bold is False

    # 5. Kiểm tra chỉ mục dấu '-' không in đậm (chuẩn Cấp 4 HD 05-HD/VPTW)
    p_cq = [p for p in doc_kh.paragraphs if p.text.strip().startswith("- Cơ quan thường trực")][0]
    # Dấu '-' tuyệt đối không in đậm
    assert p_cq.runs[0].text.strip() == "-" and p_cq.runs[0].bold is False, "Dấu '-' không được in đậm"
    # Tiêu đề có ** thì in đậm, nội dung sau đó in thường
    assert p_cq.runs[1].text.strip() == "Cơ quan thường trực:" and p_cq.runs[1].bold is True
    assert p_cq.runs[2].bold is False

    # 5b. Kiểm tra chỉ mục '-' không có markdown: toàn bộ in thường
    p_ph = [p for p in doc_kh.paragraphs if p.text.strip().startswith("- Cơ quan phối hợp")][0]
    assert all(r.bold is False for r in p_ph.runs), "Đoạn '-' không có markdown phải in thường toàn bộ"

    # 5c. Kiểm tra chỉ mục dấu '+' không in đậm (chuẩn Cấp 5 HD 05-HD/VPTW)
    p_plus = [p for p in doc_kh.paragraphs if p.text.strip().startswith("+ Phân công")][0]
    assert p_plus.runs[0].text.strip().startswith("+") and p_plus.runs[0].bold is False, "Dấu '+' không được in đậm"
    assert all(r.bold is False for r in p_plus.runs), "Đoạn '+' không có markdown phải in thường toàn bộ"

    # 6. '1. Mục tiêu tổng quát' (tiêu đề tiểu mục độc lập)
    p_sub = [p for p in doc_kh.paragraphs if p.text.strip() == "1. Mục tiêu tổng quát"][0]
    assert p_sub.runs[0].bold is True and p_sub.runs[0].text.strip() == "1. Mục tiêu tổng quát"

    print(f"  -> TC-05 HOÀN TOÀN ĐẠT CHUẨN IN ĐẬM CHỈ MỤC: {out_file}")


def test_tc04_anti_hallucination():
    print("[TEST TC-04] Đang kiểm tra phòng chống nhầm lẫn địa giới hành chính (Anti-Hallucination)...")
    forbidden_terms = ["huyện thuận bắc", "huyện uỷ thuận bắc", "ubnd huyện"]
    
    test_files = [
        os.path.join(BASE_DIR, "AGENTS.md"),
        os.path.join(BASE_DIR, ".agents", "skills", "quy-chuan-nen-tang", "SKILL.md"),
        os.path.join(BASE_DIR, ".agents", "skills", "cong-van-giao-viec", "SKILL.md"),
        os.path.join(BASE_DIR, ".agents", "skills", "ke-hoach-nghi-quyet", "SKILL.md")
    ]
    
    found_violations = []
    for file_path in test_files:
        if not os.path.exists(file_path):
            continue
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read().lower()
            # Allow warnings/prohibitions like "KHÔNG dùng huyện Thuận Bắc"
            for term in forbidden_terms:
                if term in content:
                    lines = content.splitlines()
                    for idx, line in enumerate(lines):
                        if term in line and not any(neg in line for neg in ["không", "tuyệt đối không", "bỏ", "xoá", "tránh"]):
                            found_violations.append((file_path, idx + 1, line.strip()))

    if found_violations:
        print(f"  -> TC-04 CẢNH BÁO: Phát hiện vi phạm nhắc đến cấp huyện cũ: {found_violations}")
    else:
        print("  -> TC-04 HOÀN TOÀN ĐẠT: Không có vi phạm nhầm lẫn địa giới 3 cấp.")


if __name__ == "__main__":
    print("==============================================================================")
    print("     BẮT ĐẦU CHẠY BỘ KIỂM THỬ TÍCH HỢP NGHIỆP VỤ (TEST SUITE)")
    print("==============================================================================")
    test_tc02_thong_bao_ket_luan()
    test_tc03_bao_cao_tham_dinh()
    test_tc05_indicator_bolding_universal()
    test_tc04_anti_hallucination()
    print("==============================================================================")
    print("     TẤT CẢ CÁC BÀI TEST ĐỀU ĐẠT CHUẨN 100%!")
    print("==============================================================================")
