import os, sys
import docx

doc_path = r"van_ban_den\2026\12_TB_KET_LUAN_KIEM_TRA_PCTNLPTC_DOI_VOI_CB_THON_SUOI_GIENG.docx"
doc = docx.Document(doc_path)
out_path = r"van_ban_den\2026\detailed_inspection.txt"

lines = []
lines.append("=== SECTIONS & MARGINS ===")
for s in doc.sections:
    lines.append(f"Page size: {s.page_width.mm:.1f} x {s.page_height.mm:.1f} mm")
    lines.append(f"Top: {s.top_margin.mm:.1f}mm, Bottom: {s.bottom_margin.mm:.1f}mm, Left: {s.left_margin.mm:.1f}mm, Right: {s.right_margin.mm:.1f}mm")

lines.append("\n=== TABLES ===")
for i, t in enumerate(doc.tables):
    lines.append(f"--- Table {i+1} ({len(t.rows)} rows x {len(t.columns)} cols) ---")
    for r_idx, r in enumerate(t.rows):
        for c_idx, c in enumerate(r.cells):
            cell_paragraphs = [p.text for p in c.paragraphs]
            lines.append(f"Cell [{r_idx},{c_idx}]: {cell_paragraphs}")

lines.append("\n=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    runs_detail = []
    for r in p.runs:
        font_name = r.font.name
        font_size = r.font.size.pt if r.font.size else "inherit"
        bold = r.bold
        italic = r.italic
        runs_detail.append(f"['{r.text}' fn={font_name} sz={font_size} b={bold} i={italic}]")
    
    first_indent = f"{p.paragraph_format.first_line_indent.mm:.1f}mm" if p.paragraph_format.first_line_indent else "None"
    left_indent = f"{p.paragraph_format.left_indent.mm:.1f}mm" if p.paragraph_format.left_indent else "None"
    space_before = f"{p.paragraph_format.space_before.pt:.1f}pt" if p.paragraph_format.space_before else "None"
    space_after = f"{p.paragraph_format.space_after.pt:.1f}pt" if p.paragraph_format.space_after else "None"
    line_spacing = f"{p.paragraph_format.line_spacing}" if p.paragraph_format.line_spacing else "None"
    align = str(p.alignment)
    
    lines.append(f"\n--- Paragraph {i:02d} ---")
    lines.append(f"Text: {p.text}")
    lines.append(f"Format: align={align}, indent={first_indent}, left={left_indent}, bef={space_before}, aft={space_after}, line_sp={line_spacing}")
    lines.append(f"Runs: {' '.join(runs_detail)}")

with open(out_path, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print(f"Written detailed inspection to {out_path}")
