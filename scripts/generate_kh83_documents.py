import os
import sys
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def remove_table_borders(table):
    tblBorders = parse_xml(r'''
        <w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    ''')
    table._tbl.tblPr.append(tblBorders)

def add_header_table(doc, so_hieu, trich_yeu, ngay_thang):
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(header_table)
    
    col1 = header_table.columns[0]
    col2 = header_table.columns[1]
    col1.width = Mm(75)
    col2.width = Mm(90)
    
    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)
    cell_left.width = Mm(75)
    cell_right.width = Mm(90)
    set_cell_margins(cell_left, top=0, bottom=0, left=0, right=70)
    set_cell_margins(cell_right, top=0, bottom=0, left=70, right=0)
    
    # Cột 1 (Trái): Tên đơn vị & Số ký hiệu - CĂN GIỮA
    p_cq_tren = cell_left.paragraphs[0]
    p_cq_tren.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cq_tren.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_cq_tren.paragraph_format.line_spacing = 1.05
    p_cq_tren.paragraph_format.space_before = Pt(0)
    p_cq_tren.paragraph_format.space_after = Pt(1)
    run = p_cq_tren.add_run("ĐẢNG BỘ TỈNH KHÁNH HOÀ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    
    p_cq_bh = cell_left.add_paragraph()
    p_cq_bh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cq_bh.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_cq_bh.paragraph_format.line_spacing = 1.05
    p_cq_bh.paragraph_format.space_before = Pt(0)
    p_cq_bh.paragraph_format.space_after = Pt(1)
    run = p_cq_bh.add_run("ĐẢNG UỶ XÃ CÔNG HẢI")
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    
    p_sao = cell_left.add_paragraph()
    p_sao.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sao.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_sao.paragraph_format.line_spacing = 1.0
    p_sao.paragraph_format.space_before = Pt(0)
    p_sao.paragraph_format.space_after = Pt(2)
    run = p_sao.add_run("*")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    
    p_so = cell_left.add_paragraph()
    p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_so.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_so.paragraph_format.line_spacing = 1.05
    p_so.paragraph_format.space_before = Pt(0)
    p_so.paragraph_format.space_after = Pt(3)
    run = p_so.add_run(so_hieu)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13.5)
    
    if trich_yeu:
        p_trichyeu = cell_left.add_paragraph()
        p_trichyeu.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_trichyeu.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p_trichyeu.paragraph_format.line_spacing = 1.05
        p_trichyeu.paragraph_format.space_before = Pt(0)
        p_trichyeu.paragraph_format.space_after = Pt(0)
        run = p_trichyeu.add_run(trich_yeu)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.italic = False
    
    # Cột 2 (Phải): Tiêu ngữ & Ngày tháng - CĂN PHẢI
    p_tieungu = cell_right.paragraphs[0]
    p_tieungu.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tieungu.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_tieungu.paragraph_format.line_spacing = 1.05
    p_tieungu.paragraph_format.space_before = Pt(0)
    p_tieungu.paragraph_format.space_after = Pt(1)
    run = p_tieungu.add_run("ĐẢNG CỘNG SẢN VIỆT NAM")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    
    p_line = cell_right.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_line.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_line.paragraph_format.line_spacing = 1.0
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(4)
    vml_xml = """<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml">
        <w:pict>
            <v:line from="0,0" to="142pt,0" strokecolor="#000000" strokeweight="1pt"/>
        </w:pict>
    </w:r>"""
    p_line._p.append(parse_xml(vml_xml))
    
    p_ngaythang = cell_right.add_paragraph()
    p_ngaythang.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ngaythang.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_ngaythang.paragraph_format.line_spacing = 1.05
    p_ngaythang.paragraph_format.space_before = Pt(0)
    p_ngaythang.paragraph_format.space_after = Pt(0)
    run = p_ngaythang.add_run(ngay_thang)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13.5)
    run.font.italic = True

def add_footer_table(doc, noi_nhan_list, tm_text, chuc_vu_text, ho_ten_text):
    footer_table = doc.add_table(rows=1, cols=2)
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(footer_table)
    
    col1_f = footer_table.columns[0]
    col2_f = footer_table.columns[1]
    col1_f.width = Mm(75)
    col2_f.width = Mm(90)
    
    cell_f_left = footer_table.cell(0, 0)
    cell_f_right = footer_table.cell(0, 1)
    cell_f_left.width = Mm(75)
    cell_f_right.width = Mm(90)
    set_cell_margins(cell_f_left, top=0, bottom=0, left=0, right=70)
    set_cell_margins(cell_f_right, top=0, bottom=0, left=70, right=0)
    
    # Nơi nhận
    p_nn_title = cell_f_left.paragraphs[0]
    p_nn_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_nn_title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_nn_title.paragraph_format.line_spacing = 1.05
    p_nn_title.paragraph_format.space_before = Pt(0)
    p_nn_title.paragraph_format.space_after = Pt(2)
    run_nn = p_nn_title.add_run("Nơi nhận:")
    run_nn.font.name = "Times New Roman"
    run_nn.font.size = Pt(14)
    run_nn.font.underline = True
    
    for item in noi_nhan_list:
        p_item = cell_f_left.add_paragraph()
        p_item.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_item.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p_item.paragraph_format.line_spacing = 1.05
        p_item.paragraph_format.space_before = Pt(0)
        p_item.paragraph_format.space_after = Pt(1)
        run_it = p_item.add_run(item)
        run_it.font.name = "Times New Roman"
        run_it.font.size = Pt(12)
        
    # Thẩm quyền ký
    p_tm = cell_f_right.paragraphs[0]
    p_tm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tm.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_tm.paragraph_format.line_spacing = 1.05
    p_tm.paragraph_format.space_before = Pt(0)
    p_tm.paragraph_format.space_after = Pt(2)
    run_tm = p_tm.add_run(tm_text)
    run_tm.font.name = "Times New Roman"
    run_tm.font.size = Pt(13.5)
    run_tm.font.bold = True
    
    p_cv = cell_f_right.add_paragraph()
    p_cv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cv.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_cv.paragraph_format.line_spacing = 1.05
    p_cv.paragraph_format.space_before = Pt(0)
    p_cv.paragraph_format.space_after = Pt(45)
    run_cv = p_cv.add_run(chuc_vu_text)
    run_cv.font.name = "Times New Roman"
    run_cv.font.size = Pt(13.5)
    run_cv.font.bold = True
    
    p_ten = cell_f_right.add_paragraph()
    p_ten.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ten.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_ten.paragraph_format.line_spacing = 1.05
    p_ten.paragraph_format.space_before = Pt(0)
    p_ten.paragraph_format.space_after = Pt(0)
    run_ten = p_ten.add_run(ho_ten_text)
    run_ten.font.name = "Times New Roman"
    run_ten.font.size = Pt(14)
    run_ten.font.bold = True

def format_body_paragraph(p, first_line=10, before=6, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.alignment = align
    p.paragraph_format.left_indent = Mm(0)
    p.paragraph_format.right_indent = Mm(0)
    p.paragraph_format.first_line_indent = Mm(first_line)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(18)

def setup_page(doc):
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    style.paragraph_format.line_spacing = Pt(18)
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(6)

def export_cong_van(output_path):
    doc = docx.Document()
    setup_page(doc)
    
    so_hieu = "Số       -CV/ĐU"
    trich_yeu = "Tham mưu triển khai thực hiện\nKế hoạch số 83-KH/TU của\nBan Thường vụ Tỉnh uỷ"
    ngay_thang = "Công Hải, ngày     tháng 6 năm 2026"
    add_header_table(doc, so_hieu, trich_yeu, ngay_thang)
    
    # Khoảng cách trước kính gửi
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)
    p_sp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_sp.paragraph_format.line_spacing = 1.0
    
    # Kính gửi
    p_kg = doc.add_paragraph()
    p_kg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kg.paragraph_format.first_line_indent = Mm(0)
    p_kg.paragraph_format.space_before = Pt(4)
    p_kg.paragraph_format.space_after = Pt(8)
    p_kg.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_kg.paragraph_format.line_spacing = Pt(18)
    r_kg = p_kg.add_run("Kính gửi: ")
    r_kg.font.italic = True
    r_kg.font.size = Pt(14)
    
    agencies = [
        "Uỷ ban nhân dân xã;",
        "Ban Xây dựng Đảng;",
        "Uỷ ban Kiểm tra Đảng uỷ;",
        "Uỷ ban Mặt trận Tổ quốc Việt Nam xã;",
        "Các chi bộ trực thuộc Đảng uỷ."
    ]
    for i, ag in enumerate(agencies):
        p_ag = doc.add_paragraph()
        format_body_paragraph(p_ag, first_line=15, before=1, after=1, align=WD_ALIGN_PARAGRAPH.LEFT)
        r = p_ag.add_run(f"- {ag}")
        r.font.size = Pt(14)
    
    # Thân bài
    p_body1 = doc.add_paragraph()
    format_body_paragraph(p_body1)
    p_body1.add_run("Thực hiện Kế hoạch số 83-KH/TU, ngày 12/6/2026 của Ban Thường vụ Tỉnh uỷ về thực hiện Nghị quyết Hội nghị lần thứ hai Ban Chấp hành Trung ương Đảng khoá XIV về tiếp tục tăng cường sự lãnh đạo của Đảng đối với công tác phòng, chống tham nhũng, lãng phí, tiêu cực trong giai đoạn mới. Thường trực Đảng uỷ có ý kiến chỉ đạo như sau:")
    
    # Điểm 1: Ban Xây dựng Đảng
    p_d1 = doc.add_paragraph()
    format_body_paragraph(p_d1)
    r = p_d1.add_run("1. Giao Ban Xây dựng Đảng:")
    r.font.bold = True
    
    p_d1_1 = doc.add_paragraph()
    format_body_paragraph(p_d1_1, first_line=15)
    p_d1_1.add_run("- Chủ trì tham mưu Thường trực Đảng uỷ kế hoạch và tổ chức học tập, nghiên cứu, quán triệt, tuyên truyền Nghị quyết số 04-NQ/TW của Ban Chấp hành Trung ương Đảng, Kế hoạch số 03-KH/TW của Bộ Chính trị và Kế hoạch số 83-KH/TU của Ban Thường vụ Tỉnh uỷ đến toàn thể cán bộ, đảng viên và Nhân dân; hoàn thành ")
    r = p_d1_1.add_run("trong tháng 6/2026")
    r.font.bold = True
    p_d1_1.add_run(".")
    
    p_d1_2 = doc.add_paragraph()
    format_body_paragraph(p_d1_2, first_line=15)
    p_d1_2.add_run("- Phối hợp với Văn phòng Đảng uỷ tham mưu Ban Thường vụ Đảng uỷ xây dựng Kế hoạch của Đảng uỷ xã thực hiện Kế hoạch số 83-KH/TU của Ban Thường vụ Tỉnh uỷ; trình Thường trực Đảng uỷ ")
    r = p_d1_2.add_run("trước ngày 24/6/2026")
    r.font.bold = True
    p_d1_2.add_run(".")
    
    p_d1_3 = doc.add_paragraph()
    format_body_paragraph(p_d1_3, first_line=15)
    p_d1_3.add_run("- Thường xuyên theo dõi, nắm bắt tình hình hoạt động, tư tưởng của đội ngũ cán bộ, công chức cơ sở; đẩy mạnh giáo dục văn hoá liêm chính, không tham nhũng, lãng phí; tham mưu thực hiện nghiêm cơ chế bảo vệ cán bộ năng động, sáng tạo, dám nghĩ, dám làm vì lợi ích chung; kịp thời đề xuất thay thế cán bộ có biểu hiện đùn đẩy, né tránh, sợ trách nhiệm.")
    
    # Điểm 2: UBND xã
    p_d2 = doc.add_paragraph()
    format_body_paragraph(p_d2)
    r = p_d2.add_run("2. Giao Uỷ ban nhân dân xã:")
    r.font.bold = True
    
    p_d2_1 = doc.add_paragraph()
    format_body_paragraph(p_d2_1, first_line=15)
    p_d2_1.add_run("- Căn cứ chức năng quản lý nhà nước và tình hình thực tế địa phương, tham mưu Kế hoạch của UBND xã triển khai thực hiện nhiệm vụ phòng, chống tham nhũng, lãng phí, tiêu cực trên các lĩnh vực trọng điểm: quản lý đất đai, trật tự xây dựng, tài nguyên khoáng sản, thu - chi ngân sách và quản lý tài sản công.")
    
    p_d2_2 = doc.add_paragraph()
    format_body_paragraph(p_d2_2, first_line=15)
    p_d2_2.add_run("- Khẩn trương rà soát, xử lý dứt điểm các công trình, dự án đầu tư công chậm tiến độ, tồn đọng kéo dài trên địa bàn xã; hoàn thiện phương án quản lý, sử dụng đúng mục đích, hiệu quả các cơ sở nhà đất dôi dư sau sắp xếp tổ chức bộ máy, đơn vị hành chính; báo cáo Thường trực Đảng uỷ ")
    r = p_d2_2.add_run("trước ngày 25/6/2026")
    r.font.bold = True
    p_d2_2.add_run(".")
    
    p_d2_3 = doc.add_paragraph()
    format_body_paragraph(p_d2_3, first_line=15)
    p_d2_3.add_run("- Đẩy mạnh cải cách hành chính, chuyển đổi số, công khai minh bạch quy trình thủ tục tại bộ phận Một cửa; kiên quyết ngăn chặn, xử lý tình trạng nhũng nhiễu, phiền hà, khắc phục triệt để \"tham nhũng vặt\".")
    
    # Điểm 3: UBKT Đảng uỷ
    p_d3 = doc.add_paragraph()
    format_body_paragraph(p_d3)
    r = p_d3.add_run("3. Giao Cơ quan Uỷ ban Kiểm tra Đảng uỷ:")
    r.font.bold = True
    
    p_d3_1 = doc.add_paragraph()
    format_body_paragraph(p_d3_1, first_line=15)
    p_d3_1.add_run("- Chủ trì tham mưu Kế hoạch kiểm tra, giám sát chuyên đề của Cấp uỷ và UBKT Đảng uỷ năm 2026 về công tác phòng, chống tham nhũng, lãng phí, tiêu cực; tập trung kiểm tra, giám sát trách nhiệm người đứng đầu và cán bộ, đảng viên trong các lĩnh vực dễ phát sinh sai phạm.")
    
    p_d3_2 = doc.add_paragraph()
    format_body_paragraph(p_d3_2, first_line=15)
    p_d3_2.add_run("- Tăng cường kiểm soát việc kê khai tài sản, thu nhập của cán bộ diện Ban Thường vụ Đảng uỷ quản lý; chủ động phòng ngừa, phát hiện và kiểm tra khi có dấu hiệu vi phạm; định kỳ báo cáo tiến độ về Thường trực Đảng uỷ ")
    r = p_d3_2.add_run("trước ngày 23 hàng tháng")
    r.font.bold = True
    p_d3_2.add_run(".")
    
    # Điểm 4: UBMTTQVN xã
    p_d4 = doc.add_paragraph()
    format_body_paragraph(p_d4)
    r = p_d4.add_run("4. Giao Cơ quan Uỷ ban Mặt trận Tổ quốc Việt Nam xã:")
    r.font.bold = True
    
    p_d4_1 = doc.add_paragraph()
    format_body_paragraph(p_d4_1, first_line=15)
    p_d4_1.add_run("- Chủ trì phối hợp với các tổ chức chính trị - xã hội phát huy vai trò giám sát và phản biện xã hội của Mặt trận Tổ quốc và Nhân dân; tập trung giám sát việc thực thi công vụ, quản lý đất đai, đầu tư công tại cộng đồng (thông qua Ban Giám sát đầu tư của cộng đồng và Ban Thanh tra nhân dân).")
    
    p_d4_2 = doc.add_paragraph()
    format_body_paragraph(p_d4_2, first_line=15)
    p_d4_2.add_run("- Kịp thời nắm bắt dư luận xã hội, tạo điều kiện thuận lợi để người dân, doanh nghiệp phản ánh, kiến nghị, tố giác hành vi tham nhũng, lãng phí, tiêu cực; bảo vệ người phản ánh, tố giác theo quy định.")
    
    # Điểm 5: Chi bộ trực thuộc
    p_d5 = doc.add_paragraph()
    format_body_paragraph(p_d5)
    r = p_d5.add_run("5. Các chi bộ trực thuộc Đảng uỷ:")
    r.font.bold = True
    
    p_d5_1 = doc.add_paragraph()
    format_body_paragraph(p_d5_1, first_line=15)
    p_d5_1.add_run("- Tổ chức phổ biến, quán triệt nghiêm túc Nghị quyết số 04-NQ/TW của Ban Chấp hành Trung ương Đảng và Kế hoạch số 83-KH/TU của Ban Thường vụ Tỉnh uỷ đến toàn thể đảng viên trong kỳ sinh hoạt chi bộ định kỳ tháng 6/2026; đưa nội dung này vào sinh hoạt thường xuyên gắn với tự phê bình và phê bình.")
    
    # Điểm 6: Văn phòng Đảng uỷ
    p_d6 = doc.add_paragraph()
    format_body_paragraph(p_d6)
    r = p_d6.add_run("6. Giao Văn phòng Đảng uỷ:")
    r.font.bold = True
    
    p_d6_1 = doc.add_paragraph()
    format_body_paragraph(p_d6_1, first_line=15)
    p_d6_1.add_run("- Theo dõi, đôn đốc các cơ quan, đơn vị thực hiện đúng tiến độ được giao; chủ trì thẩm định hồ sơ dự thảo Kế hoạch của Ban Thường vụ Đảng uỷ xã; tham mưu chế độ thông tin, báo cáo định kỳ gửi Tỉnh uỷ theo quy định.")
    
    # Đoạn kết & Đính kèm
    p_end = doc.add_paragraph()
    format_body_paragraph(p_end)
    p_end.add_run("Đề nghị các cơ quan, đơn vị nghiêm túc triển khai thực hiện./.")
    
    p_attach = doc.add_paragraph()
    format_body_paragraph(p_attach, first_line=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p_attach.add_run("*(Đính kèm Kế hoạch số 83-KH/TU, ngày 12/6/2026 của Ban Thường vụ Tỉnh uỷ)./.*")
    r.font.italic = True
    
    # Footer table
    noi_nhan = [
        "- Như trên;",
        "- Thường trực Đảng uỷ;",
        "- Lưu VPĐU."
    ]
    add_footer_table(doc, noi_nhan, "T/M BAN THƯỜNG VỤ", "PHÓ BÍ THƯ", "Nguyễn Xuân Hoàng")
    
    doc.save(output_path)
    print("Exported Cong Van successfully to:", output_path)

def export_ke_hoach(output_path):
    doc = docx.Document()
    setup_page(doc)
    
    so_hieu = "Số       -KH/ĐU"
    ngay_thang = "Công Hải, ngày     tháng 6 năm 2026"
    add_header_table(doc, so_hieu, None, ngay_thang)
    
    # Tiêu đề
    p_title1 = doc.add_paragraph()
    p_title1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title1.paragraph_format.first_line_indent = Mm(0)
    p_title1.paragraph_format.space_before = Pt(14)
    p_title1.paragraph_format.space_after = Pt(4)
    p_title1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_title1.paragraph_format.line_spacing = Pt(20)
    r1 = p_title1.add_run("KẾ HOẠCH")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(16)
    r1.font.bold = True
    
    p_title2 = doc.add_paragraph()
    p_title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title2.paragraph_format.first_line_indent = Mm(0)
    p_title2.paragraph_format.space_before = Pt(2)
    p_title2.paragraph_format.space_after = Pt(10)
    p_title2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_title2.paragraph_format.line_spacing = Pt(18)
    r2 = p_title2.add_run("thực hiện Kế hoạch số 83-KH/TU của Ban Thường vụ Tỉnh uỷ và Nghị quyết Hội nghị lần thứ hai Ban Chấp hành Trung ương Đảng khoá XIV về tiếp tục tăng cường sự lãnh đạo của Đảng đối với công tác phòng, chống tham nhũng, lãng phí, tiêu cực trong giai đoạn mới")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)
    r2.font.bold = True
    
    # Căn cứ
    cancus = [
        "Thực hiện Nghị quyết số 04-NQ/TW, ngày 01/4/2026 của Ban Chấp hành Trung ương Đảng khoá XIV về tiếp tục tăng cường sự lãnh đạo của Đảng đối với công tác phòng, chống tham nhũng, lãng phí, tiêu cực trong giai đoạn mới (sau đây viết tắt là Nghị quyết số 04-NQ/TW);",
        "Thực hiện Kế hoạch số 03-KH/TW, ngày 18/5/2026 của Bộ Chính trị về thực hiện Nghị quyết số 04-NQ/TW của Ban Chấp hành Trung ương Đảng;",
        "Thực hiện Kế hoạch số 83-KH/TU, ngày 12/6/2026 của Ban Thường vụ Tỉnh uỷ Khánh Hoà về thực hiện Nghị quyết Hội nghị lần thứ hai Ban Chấp hành Trung ương Đảng khoá XIV;",
        "Căn cứ Quy chế làm việc số 01-QC/ĐU, ngày 01/7/2025 của Ban Chấp hành Đảng bộ xã Công Hải khoá I, nhiệm kỳ 2025 - 2030;",
        "Ban Thường vụ Đảng uỷ xã ban hành Kế hoạch triển khai thực hiện như sau:"
    ]
    for cc in cancus:
        p_cc = doc.add_paragraph()
        format_body_paragraph(p_cc, first_line=10, before=3, after=3)
        r = p_cc.add_run(cc)
        r.font.size = Pt(14)
        if cc.endswith(":"):
            r.font.bold = False
        else:
            r.font.italic = True
            
    # I. MỤC ĐÍCH, YÊU CẦU
    p_i = doc.add_paragraph()
    format_body_paragraph(p_i, first_line=0, before=10, after=4)
    r = p_i.add_run("I. MỤC ĐÍCH, YÊU CẦU")
    r.font.bold = True
    
    p_i1 = doc.add_paragraph()
    format_body_paragraph(p_i1, first_line=10, before=4, after=2)
    r = p_i1.add_run("1. Mục đích")
    r.font.bold = True
    
    items_md = [
        "Cụ thể hoá và triển khai thực hiện nghiêm túc, hiệu quả, thực chất các quan điểm, mục tiêu, nhiệm vụ, giải pháp nêu tại Nghị quyết số 04-NQ/TW của Trung ương và Kế hoạch số 83-KH/TU của Ban Thường vụ Tỉnh uỷ phù hợp với đặc điểm, tình hình thực tiễn của xã Công Hải.",
        "Nâng cao năng lực lãnh đạo, sức chiến đấu của Đảng bộ xã; xây dựng hệ thống chính trị cơ sở trong sạch, vững mạnh, hoạt động hiệu lực, hiệu quả; kiên quyết ngăn chặn, đẩy lùi, xử lý nghiêm cán bộ, đảng viên suy thoái về tư tưởng chính trị, đạo đức, lối sống, có hành vi tham nhũng, lãng phí, tiêu cực.",
        "Tạo môi trường thông thoáng, minh bạch, giữ vững kỷ luật, kỷ cương, củng cố niềm tin của cán bộ, đảng viên và Nhân dân vào sự lãnh đạo của Đảng; góp phần thực hiện thắng lợi Nghị quyết Đại hội Đảng bộ xã nhiệm kỳ 2025 - 2030 và mục tiêu phát triển kinh tế - xã hội của tỉnh Khánh Hoà."
    ]
    for it in items_md:
        p_it = doc.add_paragraph()
        format_body_paragraph(p_it, first_line=10)
        p_it.add_run(f"- {it}")
        
    p_i2 = doc.add_paragraph()
    format_body_paragraph(p_i2, first_line=10, before=4, after=2)
    r = p_i2.add_run("2. Yêu cầu")
    r.font.bold = True
    
    items_yc = [
        "Các cấp uỷ, chi bộ, chính quyền, Mặt trận Tổ quốc và các đoàn thể chính trị - xã hội xã xác định công tác phòng, chống tham nhũng, lãng phí, tiêu cực là nhiệm vụ trọng tâm, thường xuyên, lâu dài, gắn liền với nhiệm vụ chính trị của địa phương.",
        "Thực hiện nghiêm nguyên tắc \"5 Rõ\": Rõ người chủ trì, rõ việc, rõ tiến độ thời hạn hoàn thành, rõ sản phẩm đầu ra và rõ trách nhiệm kiểm tra giám sát; khắc phục triệt để bệnh hình thức, né tránh, đùn đẩy trách nhiệm.",
        "Kết hợp chặt chẽ giữa \"xây\" và \"chống\", lấy phòng ngừa là chính; xử lý nghiêm minh, đồng bộ, kịp thời các hành vi vi phạm, không có vùng cấm, không có ngoại lệ; đồng thời bảo vệ cán bộ năng động, sáng tạo, dám nghĩ, dám làm vì lợi ích chung."
    ]
    for it in items_yc:
        p_it = doc.add_paragraph()
        format_body_paragraph(p_it, first_line=10)
        p_it.add_run(f"- {it}")

    # II. MỤC TIÊU
    p_ii = doc.add_paragraph()
    format_body_paragraph(p_ii, first_line=0, before=10, after=4)
    r = p_ii.add_run("II. MỤC TIÊU")
    r.font.bold = True
    
    p_ii1 = doc.add_paragraph()
    format_body_paragraph(p_ii1, first_line=10, before=4, after=2)
    r = p_ii1.add_run("1. Mục tiêu chung")
    r.font.bold = True
    
    p_mtc = doc.add_paragraph()
    format_body_paragraph(p_mtc, first_line=10)
    p_mtc.add_run("Tăng cường sự lãnh đạo toàn diện của Đảng bộ đối với công tác phòng, chống tham nhũng, lãng phí, tiêu cực; xây dựng văn hoá liêm chính, tiết kiệm trong toàn thể cán bộ, đảng viên, công chức; quản lý, sử dụng chặt chẽ, hiệu quả các nguồn lực đất đai, ngân sách, tài sản công; chấm dứt tình trạng nhũng nhiễu, phiền hà đối với người dân và doanh nghiệp, giữ vững ổn định chính trị - xã hội tại địa bàn xã.")
    
    p_ii2 = doc.add_paragraph()
    format_body_paragraph(p_ii2, first_line=10, before=4, after=2)
    r = p_ii2.add_run("2. Mục tiêu cụ thể")
    r.font.bold = True
    
    p_mtt = doc.add_paragraph()
    format_body_paragraph(p_mtt, first_line=10)
    r = p_mtt.add_run("a) Mục tiêu đến hết năm 2026:")
    r.font.bold = True
    
    items_2026 = [
        "100% cấp uỷ, chi bộ trực thuộc, cán bộ, đảng viên và công chức xã được quán triệt, học tập đầy đủ Nghị quyết số 04-NQ/TW của Trung ương và Kế hoạch số 83-KH/TU của Tỉnh uỷ hoàn thành trong Quý II/2026.",
        "100% cán bộ, công chức thuộc diện kê khai tài sản, thu nhập thực hiện nghiêm túc việc kê khai, công khai và minh bạch theo quy định.",
        "Hoàn thành rà soát 100% các công trình, dự án đầu tư công trên địa bàn xã; xử lý dứt điểm các vướng mắc, tồn đọng, không để phát sinh lãng phí, thất thoát nguồn lực.",
        "Hoàn thành phương án sắp xếp, quản lý và sử dụng đúng quy định đối với toàn bộ các cơ sở nhà đất, tài sản công dôi dư sau khi triển khai mô hình chính quyền địa phương mới trước ngày 30/6/2026.",
        "Tỷ lệ hồ sơ thủ tục hành chính giải quyết đúng hạn và trước hạn tại bộ phận Một cửa đạt từ 98% trở lên; 100% phản ánh, kiến nghị của người dân về hành vi nhũng nhiễu được tiếp nhận, xử lý kịp thời."
    ]
    for it in items_2026:
        p_it = doc.add_paragraph()
        format_body_paragraph(p_it, first_line=15)
        p_it.add_run(f"- {it}")
        
    p_mt_dai = doc.add_paragraph()
    format_body_paragraph(p_mt_dai, first_line=10)
    r = p_mt_dai.add_run("b) Mục tiêu giai đoạn 2026 - 2030:")
    r.font.bold = True
    
    items_2030 = [
        "Duy trì Đảng bộ xã và 100% chi bộ trực thuộc hoàn thành tốt nhiệm vụ trở lên; không có tổ chức đảng, cán bộ, đảng viên bị kỷ luật do tham nhũng, lãng phí, tiêu cực.",
        "100% các quy trình công tác, định mức chi tiêu công, quy chế quản lý tài chính, tài sản công được rà soát, sửa đổi, bổ sung và thực hiện công khai, minh bạch.",
        "Đẩy mạnh chuyển đổi số trong quản lý điều hành của Đảng uỷ và UBND xã; thực hiện dịch vụ công trực tuyến toàn trình đối với các thủ tục hành chính đủ điều kiện, ngăn ngừa tận gốc các biểu hiện tiêu cực, \"tham nhũng vặt\"."
    ]
    for it in items_2030:
        p_it = doc.add_paragraph()
        format_body_paragraph(p_it, first_line=15)
        p_it.add_run(f"- {it}")

    # III. NHIỆM VỤ VÀ GIẢI PHÁP TRỌNG TÂM
    p_iii = doc.add_paragraph()
    format_body_paragraph(p_iii, first_line=0, before=10, after=4)
    r = p_iii.add_run("III. NHIỆM VỤ VÀ GIẢI PHÁP TRỌNG TÂM")
    r.font.bold = True
    
    tasks = [
        ("1. Nâng cao trách nhiệm của cấp uỷ, người đứng đầu; siết chặt kỷ luật, kỷ cương trong thực thi công vụ",
         "Cấp uỷ, chính quyền, trước hết là người đứng đầu cấp uỷ và chính quyền xã phải thực sự gương mẫu, đi đầu trong công tác phòng, chống tham nhũng, lãng phí, tiêu cực. Chịu trách nhiệm trực tiếp, toàn diện nếu để xảy ra tham nhũng, lãng phí tại cơ quan, đơn vị, địa bàn phụ trách.\n"
         "Thực hiện nghiêm túc chế độ tiếp công dân định kỳ của Bí thư Đảng uỷ xã (ít nhất 01 ngày/tháng) và Chủ tịch UBND xã; tăng cường đối thoại trực tiếp với Nhân dân, giải quyết dứt điểm các đơn thư khiếu nại, tố cáo, phản ánh ngay từ cơ sở, không để hình thành điểm nóng.\n"
         "Khắc phục triệt để tình trạng đùn đẩy, né tránh, sợ sai, sợ trách nhiệm; xử lý nghiêm cán bộ có biểu hiện trì trệ, gây khó khăn cho người dân và tổ chức. Đồng thời, triển khai thực hiện nghiêm túc các chủ trương của Đảng về bảo vệ cán bộ năng động, sáng tạo, dám nghĩ, dám làm vì lợi ích chung."),
        
        ("2. Tăng cường tuyên truyền, giáo dục chính trị tư tưởng, xây dựng văn hoá liêm chính, tiết kiệm",
         "Đẩy mạnh công tác giáo dục chính trị, tư tưởng, đạo đức công vụ; nâng cao nhận thức, bản lĩnh chính trị cho đội ngũ cán bộ, đảng viên, công chức xã. Đưa nội dung phòng, chống tham nhũng, lãng phí, tiêu cực vào chương trình sinh hoạt chi bộ định kỳ, sinh hoạt chuyên đề hằng quý.\n"
         "Gắn kết chặt chẽ công tác phòng, chống tham nhũng, lãng phí với việc đẩy mạnh học tập và làm theo tư tưởng, đạo đức, phong cách Hồ Chí Minh; xây dựng chuẩn mực đạo đức cách mạng của cán bộ, đảng viên trong giai đoạn mới.\n"
         "Kịp thời biểu dương, nhân rộng các điển hình tiên tiến, mô hình hiệu quả trong thực hành tiết kiệm, chống lãng phí; đẩy mạnh tuyên truyền gương người tốt, việc tốt, tạo sự lan toả tích cực trong xã hội."),
        
        ("3. Đẩy mạnh phòng, chống lãng phí trong quản lý đất đai, tài nguyên, đầu tư công và tài sản công",
         "Tập trung rà soát toàn bộ các công trình, dự án đầu tư công trên địa bàn xã; theo dõi sát tiến độ thi công, giải ngân vốn đầu tư công, xử lý dứt điểm các dự án chậm tiến độ, không để kéo dài gây lãng phí nguồn lực ngân sách.\n"
         "Nâng cao hiệu quả quản lý, sử dụng đất đai, tài nguyên, khoáng sản; xử lý nghiêm các trường hợp lấn chiếm đất công, xây dựng trái phép, sử dụng đất sai mục đích trên địa bàn xã.\n"
         "Quản lý, sử dụng nghiêm ngặt tài sản công; rà soát, quản lý chặt chẽ quỹ nhà đất, tài sản dôi dư sau sắp xếp tổ chức bộ máy, đơn vị hành chính; tuyệt đối không để xảy ra thất thoát, lấn chiếm hoặc bỏ hoang gây lãng phí."),
        
        ("4. Thực hiện nghiêm các quy định về kiểm soát quyền lực, cải cách hành chính và chuyển đổi số",
         "Triển khai đồng bộ, nghiêm túc các quy định của Bộ Chính trị, Ban Bí thư về kiểm soát quyền lực, phòng, chống tham nhũng, tiêu cực trong công tác cán bộ, quản lý tài chính, tài sản công và kiểm tra, giám sát.\n"
         "Rà soát, hoàn thiện các quy chế, quy trình nội bộ, tiêu chuẩn, định mức chi tiêu công tại cơ quan Đảng uỷ và UBND xã; đảm bảo 100% hoạt động quản lý tài chính, mua sắm tài sản công được công khai, minh bạch.\n"
         "Đẩy mạnh cải cách thủ tục hành chính, số hoá hồ sơ, quy trình giải quyết công việc tại bộ phận Một cửa; nâng cao hiệu quả ứng dụng công nghệ thông tin và chuyển đổi số nhằm công khai minh bạch toàn bộ quy trình, khắc phục triệt để hiện tượng \"tham nhũng vặt\" và sách nhiễu."),
        
        ("5. Tăng cường công tác kiểm tra, giám sát của Đảng; kịp thời phát hiện, xử lý nghiêm minh các sai phạm",
         "Uỷ ban Kiểm tra Đảng uỷ chủ động xây dựng và triển khai kế hoạch kiểm tra, giám sát chuyên đề đối với các chi bộ và đảng viên trong thực hiện các quy định về phòng, chống tham nhũng, lãng phí, tiêu cực; trọng tâm là các vị trí, lĩnh vực dễ phát sinh sai phạm như tài chính, đất đai, đầu tư công.\n"
         "Thực hiện nghiêm túc công tác kiểm soát tài sản, thu nhập của cán bộ diện Ban Thường vụ Đảng uỷ quản lý; chủ động nắm bắt thông tin, kịp thời kiểm tra khi có dấu hiệu vi phạm để xử lý từ sớm, từ xa, không để vi phạm nhỏ tích tụ thành sai phạm lớn.\n"
         "Theo dõi, đôn đốc và giám sát việc thực hiện các kết luận sau thanh tra, kiểm tra, kiểm toán của cấp trên; kiên quyết khắc phục triệt để các hạn chế, thiếu sót đã được chỉ ra."),
        
        ("6. Phát huy vai trò giám sát, phản biện xã hội của Mặt trận Tổ quốc, các đoàn thể chính trị - xã hội và Nhân dân",
         "Mặt trận Tổ quốc và các tổ chức chính trị - xã hội xã phát huy hiệu quả vai trò giám sát và phản biện xã hội theo Quyết định số 217-QĐ/TW, Quyết định số 218-QĐ/TW của Bộ Chính trị; tăng cường giám sát việc tu dưỡng, rèn luyện đạo đức, lối sống của cán bộ, đảng viên tại nơi làm việc và nơi cư trú.\n"
         "Nâng cao hiệu quả hoạt động của Ban Thanh tra nhân dân và Ban Giám sát đầu tư của cộng đồng đối với các công trình, dự án xây dựng cơ bản, các chương trình mục tiêu quốc gia trên địa bàn xã.\n"
         "Tạo điều kiện thuận lợi và cơ chế an toàn để Nhân dân, doanh nghiệp tích cực cung cấp thông tin, phản ánh, tố giác các hành vi tiêu cực, tham nhũng, lãng phí; bảo vệ quyền và lợi ích hợp pháp của người phản ánh, tố giác theo đúng quy định của Đảng và Nhà nước.")
    ]
    
    for title, content in tasks:
        p_t = doc.add_paragraph()
        format_body_paragraph(p_t, first_line=10, before=6, after=2)
        r = p_t.add_run(title)
        r.font.bold = True
        
        for para in content.split("\n"):
            p_c = doc.add_paragraph()
            format_body_paragraph(p_c, first_line=15, before=2, after=2)
            p_c.add_run(para)

    # IV. KINH PHÍ THỰC HIỆN
    p_iv = doc.add_paragraph()
    format_body_paragraph(p_iv, first_line=0, before=10, after=4)
    r = p_iv.add_run("IV. KINH PHÍ THỰC HIỆN")
    r.font.bold = True
    
    p_kp = doc.add_paragraph()
    format_body_paragraph(p_kp, first_line=10)
    p_kp.add_run("Kinh phí thực hiện Kế hoạch được bố trí từ nguồn ngân sách nhà nước theo phân cấp ngân sách hiện hành, kinh phí hoạt động của Đảng bộ xã và các nguồn kinh phí hợp pháp khác theo đúng quy định của pháp luật; bảo đảm quản lý, sử dụng đúng mục đích, tiết kiệm và hiệu quả.")

    # V. TỔ CHỨC THỰC HIỆN
    p_v = doc.add_paragraph()
    format_body_paragraph(p_v, first_line=0, before=10, after=4)
    r = p_v.add_run("V. TỔ CHỨC THỰC HIỆN")
    r.font.bold = True
    
    org_tasks = [
        ("1. Uỷ ban nhân dân xã",
         "Xây dựng Kế hoạch của UBND xã để cụ thể hoá các nhiệm vụ, giải pháp quản lý nhà nước về phòng, chống tham nhũng, lãng phí, tiêu cực trên địa bàn; hoàn thành trước ngày 30/6/2026.\n"
         "Tập trung chỉ đạo rà soát, xử lý các dự án đầu tư công chậm tiến độ; lập phương án sắp xếp, quản lý quỹ tài sản công, nhà đất dôi dư sau sắp xếp; tăng cường quản lý đất đai, ngân sách; siết chặt kỷ luật, kỷ cương hành chính và nâng cao hiệu quả hoạt động của bộ phận Một cửa."),
        
        ("2. Ban Xây dựng Đảng",
         "Chủ trì tham mưu Thường trực Đảng uỷ kế hoạch và tổ chức học tập, nghiên cứu, quán triệt, tuyên truyền Nghị quyết số 04-NQ/TW của Ban Chấp hành Trung ương Đảng và Kế hoạch số 83-KH/TU của Tỉnh uỷ; hoàn thành trong tháng 6/2026.\n"
         "Chủ trì theo dõi, rà soát, đánh giá tình hình cán bộ, công chức cơ sở; tham mưu thực hiện cơ chế bảo vệ cán bộ dám nghĩ, dám làm; tham mưu xử lý, thay thế cán bộ năng lực hạn chế, sợ trách nhiệm, né tránh."),
        
        ("3. Cơ quan Uỷ ban Kiểm tra Đảng uỷ xã",
         "Tham mưu Ban Thường vụ Đảng uỷ kế hoạch kiểm tra, giám sát chuyên đề đối với các chi bộ, cán bộ, đảng viên trong việc thực hiện nhiệm vụ phòng, chống tham nhũng, lãng phí, tiêu cực.\n"
         "Thực hiện nghiêm công tác kiểm soát tài sản, thu nhập; kịp thời phát hiện, kiểm tra khi có dấu hiệu vi phạm; xử lý nghiêm minh các sai phạm theo thẩm quyền."),
        
        ("4. Cơ quan Uỷ ban Mặt trận Tổ quốc Việt Nam xã",
         "Chủ trì phối hợp với các tổ chức chính trị - xã hội xây dựng kế hoạch giám sát, phản biện xã hội; đẩy mạnh tuyên truyền, vận động đoàn viên, hội viên và Nhân dân tích cực tham gia phòng, chống tham nhũng, lãng phí, tiêu cực.\n"
         "Phát huy vai trò của Ban Thanh tra nhân dân, Ban Giám sát đầu tư của cộng đồng; thường xuyên nắm bắt dư luận xã hội, phản ánh về Đảng uỷ và UBND xã."),
        
        ("5. Văn phòng Đảng uỷ xã",
         "Là đầu mối thường trực giúp Ban Thường vụ Đảng uỷ theo dõi, đôn đốc, kiểm tra việc triển khai thực hiện Kế hoạch này.\n"
         "Thực hiện nghiêm quy định về thẩm định thể thức, nội dung các văn bản; chủ trì phối hợp với các ban Đảng tham mưu báo cáo định kỳ gửi Tỉnh uỷ (qua Ban Nội chính Tỉnh uỷ) theo quy định."),
        
        ("6. Các chi bộ trực thuộc Đảng uỷ",
         "Căn cứ Kế hoạch này, tổ chức quán triệt sâu rộng đến toàn thể cán bộ, đảng viên trong sinh hoạt định kỳ tháng 6/2026; xây dựng kế hoạch thực hiện của chi bộ phù hợp với nhiệm vụ thực tế của thôn, cơ quan, đơn vị.\n"
         "Nâng cao chất lượng tự phê bình và phê bình, tăng cường quản lý, giám sát cán bộ, đảng viên; kịp thời phát hiện, chấn chỉnh các biểu hiện tiêu cực, lãng phí ngay từ chi bộ."),
        
        ("7. Chế độ thông tin, báo cáo",
         "Các cơ quan, ban ngành, đoàn thể xã và các chi bộ trực thuộc định kỳ báo cáo kết quả thực hiện Kế hoạch về Ban Thường vụ Đảng uỷ (qua Văn phòng Đảng uỷ) trước ngày 20 của tháng cuối quý và trước ngày 15/11 hằng năm để tổng hợp, báo cáo Tỉnh uỷ./.")
    ]
    
    for title, content in org_tasks:
        p_t = doc.add_paragraph()
        format_body_paragraph(p_t, first_line=10, before=6, after=2)
        r = p_t.add_run(title)
        r.font.bold = True
        
        for para in content.split("\n"):
            p_c = doc.add_paragraph()
            format_body_paragraph(p_c, first_line=15, before=2, after=2)
            p_c.add_run(para)
            
    # Footer table
    noi_nhan = [
        "- Thường trực Tỉnh uỷ (để b/c);",
        "- Ban Nội chính Tỉnh uỷ (để b/c);",
        "- Các ban Đảng Tỉnh uỷ (để b/c);",
        "- Ban Thường vụ, BCH Đảng bộ xã;",
        "- Thường trực HĐND, UBND xã;",
        "- UBMTTQ và các tổ chức CT-XH xã;",
        "- Các chi bộ trực thuộc Đảng uỷ;",
        "- Lưu VPĐU."
    ]
    add_footer_table(doc, noi_nhan, "T/M BAN THƯỜNG VỤ", "BÍ THƯ", "Vũ Thị Thuỳ Trang")
    
    doc.save(output_path)
    print("Exported Ke Hoach successfully to:", output_path)

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    
    out_cv = os.path.join(base_dir, "van_ban_du_thao", "2026", "Cong_van", "CV_giao_viec_KH83_TU.docx")
    out_kh = os.path.join(base_dir, "van_ban_du_thao", "2026", "Ke_hoach", "KH_thuc_hien_KH83_TU_Dang_uy_Cong_Hai.docx")
    
    os.makedirs(os.path.dirname(out_cv), exist_ok=True)
    os.makedirs(os.path.dirname(out_kh), exist_ok=True)
    
    export_cong_van(out_cv)
    export_ke_hoach(out_kh)
