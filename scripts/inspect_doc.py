import os, sys
if sys.stdout.encoding != 'utf-8':
    try: sys.stdout.reconfigure(encoding='utf-8')
    except: pass

import docx

doc_path = r"van_ban_den\2026\12_TB_KET_LUAN_KIEM_TRA_PCTNLPTC_DOI_VOI_CB_THON_SUOI_GIENG.docx"
doc = docx.Document(doc_path)

print(f"Total tables: {len(doc.tables)}")
for i, table in enumerate(doc.tables):
    print(f"\n--- TABLE {i+1} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell_text = cell.text.replace('\n', ' \\n ')
            print(f"  Row {r_idx}, Col {c_idx}: {cell_text}")

print(f"\nTotal paragraphs: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        runs_info = " | ".join([f"'{r.text}'(b={r.bold}, i={r.italic}, sz={r.font.size.pt if r.font.size else None})" for r in p.runs[:4]])
        print(f"P{i:02d} [{p.alignment}]: {p.text[:80]}... (Indent: {p.paragraph_format.first_line_indent}, Spacing: Bef={p.paragraph_format.space_before}, Aft={p.paragraph_format.space_after})")
        print(f"     Runs: {runs_info}")
