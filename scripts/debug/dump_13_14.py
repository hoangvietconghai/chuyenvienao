import os, sys
if hasattr(sys.stdout, 'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')
import docx

def dump_file(docx_path):
    print(f"\n=======================================================")
    print(f"FILE: {docx_path}")
    print(f"=======================================================")
    doc = docx.Document(docx_path)
    
    print("\n--- TABLES ---")
    for t_idx, t in enumerate(doc.tables):
        print(f"Table {t_idx+1} ({len(t.rows)}x{len(t.columns)}):")
        for r_idx, r in enumerate(t.rows):
            for c_idx, c in enumerate(r.cells):
                txt = " \\n ".join([p.text.strip() for p in c.paragraphs if p.text.strip()])
                print(f"  [{r_idx},{c_idx}]: {txt}")
                
    print("\n--- PARAGRAPHS ---")
    for p_idx, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t:
            indent = f"{p.paragraph_format.first_line_indent.mm:.1f}mm" if p.paragraph_format.first_line_indent else "None"
            bef = f"{p.paragraph_format.space_before.pt:.1f}pt" if p.paragraph_format.space_before else "None"
            aft = f"{p.paragraph_format.space_after.pt:.1f}pt" if p.paragraph_format.space_after else "None"
            linesp = p.paragraph_format.line_spacing
            print(f"P{p_idx:02d} (ind={indent}, bef={bef}, aft={aft}): {t}")

dump_file(r"van_ban_den\2026\13. TB KẾT LUẬN KIỂM TRA PCTNLPTC ĐỐI VỚI CB MẪU GIÁO CH.docx")
dump_file(r"van_ban_den\2026\14. TB KẾT LUẬN KIỂM TRA PCTNLPTC ĐỐI VỚI CB CÔNG AN XÃ.docx")
