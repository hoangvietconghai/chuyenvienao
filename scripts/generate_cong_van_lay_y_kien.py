# -*- coding: utf-8 -*-
"""
scripts/generate_cong_van_lay_y_kien.py
Tạo 2 tệp văn bản Word (.docx):
1. Công văn lấy ý kiến Ủy viên Ban Thường vụ Đảng uỷ (-CV/ĐU)
2. Phiếu xin ý kiến Ủy viên Ban Thường vụ Đảng uỷ
"""

import os
import sys
from pathlib import Path

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT

try:
    from scripts.export_docx import (
        PartyDocumentBuilder,
        set_cell_margins,
        remove_table_borders,
        add_horizontal_line,
        parse_markdown_runs
    )
except ImportError:
    from export_docx import (
        PartyDocumentBuilder,
        set_cell_margins,
        remove_table_borders,
        add_horizontal_line,
        parse_markdown_runs
    )


def generate_cong_van_lay_y_kien():
    """Tạo Công văn lấy ý kiến Ủy viên Ban Thường vụ Đảng uỷ."""
    data_cv = {
        "doc_type": "CV",
        "so_hieu": "Số        -CV/ĐU",
        "co_quan_cap_tren": "ĐẢNG BỘ TỈNH KHÁNH HOÀ",
        "co_quan_ban_hanh": "ĐẢNG UỶ XÃ CÔNG HẢI",
        "dia_danh_ngay": "Công Hải, ngày   tháng   năm 2026",
        "trich_yeu": "Góp ý Dự thảo Thông báo Kết luận kiểm tra của\nBan Thường vụ Đảng uỷ",
        "kinh_gui": ["Các đồng chí Uỷ viên Ban Thường vụ Đảng uỷ."],
        "noi_dung": [
            "Tiếp nhận Báo cáo kết quả kiểm tra của Đoàn kiểm tra theo Quyết định số 56-QĐ/ĐU ngày 15/8/2026 của Ban Thường vụ Đảng uỷ và Báo cáo thẩm định số 18-BC/VPĐU ngày 02/10/2026 của Văn phòng Đảng uỷ về việc thẩm định dự thảo Thông báo Kết luận kiểm tra việc lãnh đạo, chỉ đạo triển khai thực hiện công tác phòng, chống tham nhũng, lãng phí, tiêu cực đối với Chi bộ thôn Suối Giếng;",
            "Thường trực Đảng uỷ kính gửi các đồng chí Uỷ viên Ban Thường vụ Đảng uỷ nghiên cứu, rà soát và tham gia đóng góp ý kiến về dự thảo Thông báo Kết luận kiểm tra việc lãnh đạo, chỉ đạo triển khai thực hiện công tác phòng, chống tham nhũng, lãng phí, tiêu cực đối với Chi bộ thôn Suối Giếng *(hồ sơ gửi kèm gồm: dự thảo Thông báo Kết luận đã được Văn phòng Đảng uỷ rà soát chuẩn hóa, Báo cáo thẩm định của Văn phòng Đảng uỷ và Phiếu xin ý kiến)*.",
            "Đề nghị các đồng chí nghiên cứu, tham gia đóng góp ý kiến bằng văn bản (theo mẫu Phiếu xin ý kiến gửi kèm) hoặc góp ý trực tiếp vào dự thảo; ý kiến góp ý gửi về Thường trực Đảng uỷ (qua Văn phòng Đảng uỷ) **trước ngày 08/10/2026** để tổng hợp, tiếp thu và hoàn thiện văn bản trình đồng chí Bí thư Đảng uỷ ký ban hành chính thức.",
            "Qua thời gian nêu trên, nếu các đồng chí không có ý kiến gửi về xem như thống nhất với nội dung dự thảo Thông báo Kết luận./."
        ],
        "noi_nhan": [
            "- Như trên;",
            "- Thường trực Đảng uỷ (để b/c);",
            "- Lưu VPĐU (CV-2026)."
        ],
        "tham_quyen": "T/L BAN THƯỜNG VỤ",
        "chuc_vu": "PHÓ CHÁNH VĂN PHÒNG",
        "nguoi_ky": "Ngô Hoàng Việt"
    }

    builder = PartyDocumentBuilder(data_cv)
    builder.build_header()
    builder.build_title_section()
    builder.build_body()
    builder.build_footer()

    out_file = str(BASE_DIR / "van_ban_du_thao" / "CV_lay_y_kien_BTV_ve_TBKL_kiem_tra_Chi_bo_Suoi_Gieng.docx")
    builder.save(out_file)
    print(f"Đã xuất Công văn lấy ý kiến: {out_file}")
    return out_file


def generate_phieu_xin_y_kien():
    """Tạo tệp Phiếu xin ý kiến Ủy viên Ban Thường vụ Đảng uỷ."""
    doc = docx.Document()
    
    # Page setup A4
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)

    # Style
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    style.paragraph_format.line_spacing = Pt(18)
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(6)

    # 1. Header Table
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(tbl)

    col1 = tbl.columns[0]
    col2 = tbl.columns[1]
    col1.width = Mm(75)
    col2.width = Mm(90)
    cell_l = tbl.cell(0, 0)
    cell_r = tbl.cell(0, 1)
    cell_l.width = Mm(75)
    cell_r.width = Mm(90)
    set_cell_margins(cell_l, top=0, bottom=0, left=0, right=70)
    set_cell_margins(cell_r, top=0, bottom=0, left=70, right=0)

    # Left cell
    p1 = cell_l.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p1.paragraph_format.line_spacing = 1.05
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(1)
    r = p1.add_run("ĐẢNG BỘ TỈNH KHÁNH HOÀ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    p2 = cell_l.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p2.paragraph_format.line_spacing = 1.05
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(1)
    r = p2.add_run("ĐẢNG UỶ XÃ CÔNG HẢI")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True

    p3 = cell_l.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p3.paragraph_format.line_spacing = 1.0
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(2)
    r = p3.add_run("*")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True

    # Right cell
    p_tn = cell_r.paragraphs[0]
    p_tn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tn.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_tn.paragraph_format.line_spacing = 1.05
    p_tn.paragraph_format.space_before = Pt(0)
    p_tn.paragraph_format.space_after = Pt(1)
    r = p_tn.add_run("ĐẢNG CỘNG SẢN VIỆT NAM")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True

    p_line = cell_r.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_line.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_line.paragraph_format.line_spacing = 1.0
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(4)
    add_horizontal_line(p_line, width_pt=186, weight_pt=0.75)

    p_ngay = cell_r.add_paragraph()
    p_ngay.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ngay.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_ngay.paragraph_format.line_spacing = 1.05
    p_ngay.paragraph_format.space_before = Pt(0)
    p_ngay.paragraph_format.space_after = Pt(0)
    r = p_ngay.add_run("Công Hải, ngày   tháng   năm 2026")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13.5)
    r.font.italic = True

    # Khoảng cách
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(6)
    p_sp.paragraph_format.space_after = Pt(4)

    # 2. Tiêu đề Phiếu
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_title.paragraph_format.line_spacing = Pt(18)
    r = p_title.add_run("PHIẾU XIN Ý KIẾN")
    r.font.name = "Times New Roman"
    r.font.size = Pt(15)
    r.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    p_sub.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_sub.paragraph_format.line_spacing = Pt(18)
    r = p_sub.add_run("CỦA ĐỒNG CHÍ UỶ VIÊN BAN THƯỜNG VỤ ĐẢNG UỶ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True

    p_trich = doc.add_paragraph()
    p_trich.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_trich.paragraph_format.space_before = Pt(0)
    p_trich.paragraph_format.space_after = Pt(2)
    p_trich.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_trich.paragraph_format.line_spacing = Pt(18)
    r = p_trich.add_run("Về dự thảo Thông báo Kết luận kiểm tra việc lãnh đạo, chỉ đạo triển khai thực hiện\ncông tác phòng, chống tham nhũng, lãng phí, tiêu cực đối với Chi bộ thôn Suối Giếng")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True

    p_kem = doc.add_paragraph()
    p_kem.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kem.paragraph_format.space_before = Pt(0)
    p_kem.paragraph_format.space_after = Pt(10)
    p_kem.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_kem.paragraph_format.line_spacing = Pt(18)
    r = p_kem.add_run("(Kèm theo Công văn số        -CV/ĐU ngày   /   /2026 của Ban Thường vụ Đảng uỷ)")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.italic = True

    # 3. Kính gửi
    p_kg = doc.add_paragraph()
    p_kg.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kg.paragraph_format.first_line_indent = Mm(10)
    p_kg.paragraph_format.space_before = Pt(6)
    p_kg.paragraph_format.space_after = Pt(6)
    p_kg.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_kg.paragraph_format.line_spacing = Pt(18)
    r = p_kg.add_run("Kính gửi: ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.italic = True
    r2 = p_kg.add_run("Đồng chí ......................................................, Uỷ viên Ban Thường vụ Đảng uỷ xã.")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)

    # 4. Thân phiếu
    p_dan = doc.add_paragraph()
    p_dan.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_dan.paragraph_format.first_line_indent = Mm(10)
    p_dan.paragraph_format.space_before = Pt(6)
    p_dan.paragraph_format.space_after = Pt(6)
    p_dan.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_dan.paragraph_format.line_spacing = Pt(18)
    r = p_dan.add_run("Thực hiện Quy chế làm việc số 01-QC/ĐU của Ban Chấp hành Đảng bộ xã khoá I, nhiệm kỳ 2025 - 2030, Thường trực Đảng uỷ xin ý kiến đồng chí về dự thảo Thông báo Kết luận kiểm tra việc lãnh đạo, chỉ đạo triển khai thực hiện công tác phòng, chống tham nhũng, lãng phí, tiêu cực đối với Chi bộ thôn Suối Giếng như sau:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    # Mục 1
    p_m1 = doc.add_paragraph()
    p_m1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_m1.paragraph_format.first_line_indent = Mm(10)
    p_m1.paragraph_format.space_before = Pt(6)
    p_m1.paragraph_format.space_after = Pt(4)
    p_m1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_m1.paragraph_format.line_spacing = Pt(18)
    parse_markdown_runs(p_m1, "**1. Về toàn bộ nội dung dự thảo:**", base_font_size=14, default_bold=True)

    opt_list = [
        "[   ] Thống nhất hoàn toàn với nội dung dự thảo Thông báo Kết luận.",
        "[   ] Cơ bản thống nhất nhưng có một số ý kiến tham gia cụ thể (nêu tại Mục 2 dưới đây).",
        "[   ] Không thống nhất với dự thảo (nêu rõ lý do tại Mục 2 dưới đây)."
    ]
    for opt in opt_list:
        p_opt = doc.add_paragraph()
        p_opt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_opt.paragraph_format.first_line_indent = Mm(15)
        p_opt.paragraph_format.space_before = Pt(3)
        p_opt.paragraph_format.space_after = Pt(3)
        p_opt.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p_opt.paragraph_format.line_spacing = Pt(18)
        r = p_opt.add_run(opt)
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)

    # Mục 2
    p_m2 = doc.add_paragraph()
    p_m2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_m2.paragraph_format.first_line_indent = Mm(10)
    p_m2.paragraph_format.space_before = Pt(6)
    p_m2.paragraph_format.space_after = Pt(4)
    p_m2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_m2.paragraph_format.line_spacing = Pt(18)
    parse_markdown_runs(p_m2, "**2. Ý kiến tham gia góp ý cụ thể:**", base_font_size=14, default_bold=True)

    dots = [
        "- Về đánh giá ưu điểm, hạn chế của Chi bộ: ..........................................................................................",
        ".........................................................................................................................................................................",
        "- Về nội dung yêu cầu khắc phục và phân công trách nhiệm tổ chức thực hiện: .............................",
        ".........................................................................................................................................................................",
        "- Về phân kỳ lộ trình, thời hạn thực hiện (trước 31/10/2026 và trước 30/11/2026): ...............................",
        ".........................................................................................................................................................................",
        "- Ý kiến khác (nếu có): ......................................................................................................................................",
        "........................................................................................................................................................................."
    ]
    for dot in dots:
        p_d = doc.add_paragraph()
        p_d.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_d.paragraph_format.first_line_indent = Mm(10)
        p_d.paragraph_format.space_before = Pt(2)
        p_d.paragraph_format.space_after = Pt(2)
        p_d.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p_d.paragraph_format.line_spacing = Pt(18)
        r = p_d.add_run(dot)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)

    # Lưu ý
    p_note = doc.add_paragraph()
    p_note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_note.paragraph_format.first_line_indent = Mm(10)
    p_note.paragraph_format.space_before = Pt(8)
    p_note.paragraph_format.space_after = Pt(12)
    p_note.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_note.paragraph_format.line_spacing = Pt(18)
    r = p_note.add_run("* Phiếu xin ý kiến đề nghị gửi về Thường trực Đảng uỷ (qua Văn phòng Đảng uỷ) trước ngày 08/10/2026 để tổng hợp báo cáo Thường trực Đảng uỷ xem xét, quyết định.*")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.italic = True

    # 5. Khối Ký tên
    p_sign_t = doc.add_paragraph()
    p_sign_t.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign_t.paragraph_format.right_indent = Mm(10)
    p_sign_t.paragraph_format.space_before = Pt(6)
    p_sign_t.paragraph_format.space_after = Pt(2)
    p_sign_t.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_sign_t.paragraph_format.line_spacing = 1.05
    r = p_sign_t.add_run("UỶ VIÊN BAN THƯỜNG VỤ ĐẢNG UỶ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13.5)
    r.font.bold = True

    p_sign_sub = doc.add_paragraph()
    p_sign_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign_sub.paragraph_format.right_indent = Mm(25)
    p_sign_sub.paragraph_format.space_before = Pt(0)
    p_sign_sub.paragraph_format.space_after = Pt(0)
    p_sign_sub.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_sign_sub.paragraph_format.line_spacing = 1.05
    r = p_sign_sub.add_run("(Ký và ghi rõ họ tên)")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.italic = True

    # 5 dòng trống
    for _ in range(5):
        p_b = doc.add_paragraph()
        p_b.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p_b.paragraph_format.line_spacing = Pt(14)
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(0)

    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_name.paragraph_format.right_indent = Mm(20)
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(0)
    r = p_name.add_run("......................................................")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True

    out_file = str(BASE_DIR / "van_ban_du_thao" / "Phieu_xin_y_kien_BTV_ve_TBKL_kiem_tra_Chi_bo_Suoi_Gieng.docx")
    doc.save(out_file)
    print(f"Đã xuất Phiếu xin ý kiến: {out_file}")
    return out_file


if __name__ == "__main__":
    generate_cong_van_lay_y_kien()
    generate_phieu_xin_y_kien()
