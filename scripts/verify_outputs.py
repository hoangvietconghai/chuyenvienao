# -*- coding: utf-8 -*-
import os
import sys
from pathlib import Path

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = Path(__file__).resolve().parent.parent

import docx

def check_doc(path):
    print(f"\n=== CHECKING: {path} ===")
    d = docx.Document(path)
    s = d.sections[0]
    print(f"Margins: L={s.left_margin.mm:.1f}mm, R={s.right_margin.mm:.1f}mm, T={s.top_margin.mm:.1f}mm, B={s.bottom_margin.mm:.1f}mm")
    print(f"Tables: {len(d.tables)}")
    t1 = d.tables[0]
    print("Header Left:", [p.text for p in t1.cell(0, 0).paragraphs])
    print("Header Right:", [p.text for p in t1.cell(0, 1).paragraphs])
    t2 = d.tables[1]
    print("Footer Left (Noi nhan):", [p.text for p in t2.cell(0, 0).paragraphs[:3]])
    print("Footer Right (Ky):", [p.text for p in t2.cell(0, 1).paragraphs[:3]], "...", [p.text for p in t2.cell(0, 1).paragraphs[-2:]])
    
    for i, p in enumerate(d.paragraphs):
        if any(kw in p.text for kw in ["1. Thể thức", "2. Nội dung", "I. TỔNG QUAN", "1. Ưu điểm"]):
            first_run = p.runs[0] if p.runs else None
            fr_text = first_run.text if first_run else ""
            fr_b = first_run.bold if first_run else None
            indent = f"{p.paragraph_format.first_line_indent.mm:.1f}mm" if p.paragraph_format.first_line_indent else "None"
            bef = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else None
            aft = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else None
            linesp = p.paragraph_format.line_spacing
            print(f"  P{i:02d}: '{p.text[:45]}...' Indent={indent} Bef={bef}pt Aft={aft}pt Sp={linesp} [R0: '{fr_text}' b={fr_b}]")

if __name__ == "__main__":
    doc1 = str(BASE_DIR / "van_ban_du_thao" / "BC_tham_dinh_du_thao_TBKL_kiem_tra_Chi_bo_Suoi_Gieng.docx")
    doc2 = str(BASE_DIR / "van_ban_du_thao" / "TBKL_kiem_tra_chi_bo_thon_Suoi_Gieng_chuan_hoa.docx")
    if os.path.exists(doc1):
        check_doc(doc1)
    if os.path.exists(doc2):
        check_doc(doc2)

